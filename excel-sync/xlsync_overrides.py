# -*- coding: utf-8 -*-
"""xlsync_overrides - 코드로 옮기기 애매한 칸 단위 수정을 '덮어쓰기 목록'(JSON)으로 저장하고, 생성 코드가 저장 직전에 적용한다.

    from xlsync_overrides import apply_overrides      # 생성 코드 마지막(wb.save 직전)
    apply_overrides(wb, "overrides.json")

목록은 diff 결과에서 결정론적으로 만든다(값·수식, 서식 전체, 열너비·행높이·숨김, 병합 추가/해제).
행·열 삽입/삭제, 시트 추가·이름 변경, 체크박스·VBA·그림은 목록으로 옮기지 않는다(코드에서 직접 반영).
"""
import datetime
import json
import os
import re

from xlsync_refs import expand, rc

SUPPORTED = ("value", "style", "width", "height", "merge")


def _entry_key(e):
    t = "merge" if ("merge" in e or "unmerge" in e) else "col" if "col" in e else "row" if "row" in e else "style" if "style" in e else "value"
    tgt = e.get("merge") or e.get("unmerge") or e.get("col") or e.get("row") or " ".join(e.get("range", []))
    return (e["sheet"], t, str(tgt))


def make(result, cur, kinds=SUPPORTED):
    """diff 결과 + 현재 스냅샷 -> (항목 목록, 건너뛴 사유)."""
    structural = {c["sheet"] for c in result["changes"] if c["kind"] in ("row_insert", "row_delete", "col_insert", "col_delete")}
    sheets = {s["name"]: s for s in cur["sheets"]}
    entries, skipped = [], []
    for c in result["changes"]:
        k, sh = c["kind"], c["sheet"]
        if k not in kinds and not (k == "height" and "height" in kinds):
            if k not in SUPPORTED:
                skipped.append("%s %s %s - 코드에서 직접 반영" % (sh, k, c.get("at") or c.get("range") or c.get("attr") or ""))
            continue
        if sh in structural:
            skipped.append("%s %s %s - 같은 시트에 행·열 삽입/삭제가 있어 칸 주소가 흔들림(코드 반영 후 다시 diff)" % (sh, k, c.get("range") or c.get("at")))
            continue
        S = sheets[sh]
        if k == "value":
            cells = [addr for rng in c["range"] for addr in _addrs(rng)]
            vals = {}
            for addr in cells:
                vals.setdefault(json.dumps(S["cells"].get(addr, {}).get("v"), sort_keys=True, ensure_ascii=False), []).append(addr)
            for v, addrs in sorted(vals.items()):
                entries.append({"sheet": sh, "range": _group(addrs), "value": json.loads(v)})
        elif k == "style":
            by = {}
            for rng in c["range"]:
                for addr in _addrs(rng):
                    by.setdefault(S["cells"].get(addr, {}).get("s", S["default"]), []).append(addr)
            for s_id, addrs in sorted(by.items()):
                entries.append({"sheet": sh, "range": _group(addrs), "style": cur["styles"].get(s_id, {})})
        elif k == "width":
            e = {"sheet": sh, "col": c["at"]}
            e["hidden" if c["attr"] == "hidden" else "width"] = c["after"]
            entries.append(e)
        elif k == "height":
            e = {"sheet": sh, "row": int(c["at"])}
            e["hidden" if c["attr"] == "hidden" else "height"] = c["after"]
            entries.append(e)
        elif k == "merge":
            entries.append({"sheet": sh, ("merge" if c["attr"] == "add" else "unmerge"): c["range"][0]})
    return _dedupe(entries), skipped


def _addrs(rng):
    from xlsync_refs import a1
    return [a1(r, c) for r, c in expand(rng)]


def _group(addrs):
    from xlsync_refs import group
    return group([rc(a) for a in addrs])


def _dedupe(entries):
    out = {}
    for e in entries:
        k = _entry_key(e)
        if k[1] in ("col", "row") and k in out:
            out[k].update(e)
        else:
            out[k] = e
    return list(out.values())


def save(entries, path, note=""):
    """기존 목록에 합친다(같은 시트·종류·대상은 새 항목으로 교체). 반환: 전체 항목 수."""
    old = []
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            old = json.load(f).get("entries", [])
    merged = {_entry_key(e): e for e in old}
    for e in entries:
        merged[_entry_key(e)] = e
    data = {"xlsync_overrides": 1, "updated": datetime.datetime.now().isoformat(timespec="seconds"), "note": note,
            "entries": sorted(merged.values(), key=lambda e: (e["sheet"], _entry_key(e)[1], _entry_key(e)[2]))}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    return len(data["entries"])


# ── 적용(생성 코드에서 호출) ─────────────────────────────────────────────
def _color(s):
    from openpyxl.styles.colors import Color
    if not s or s == "auto":
        return None
    m = re.match(r"theme(\d+)([+-][\d.]+)?$", s)
    if m:
        return Color(theme=int(m.group(1)), tint=float(m.group(2) or 0))
    m = re.match(r"idx(\d+)$", s)
    if m:
        return Color(indexed=int(m.group(1)))
    return Color(rgb=("FF" + s) if len(s) == 6 else s)


def style_objects(st):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    f = st.get("font", {})
    font = Font(name=f.get("name"), sz=f.get("sz"), b=f.get("b", False), i=f.get("i", False), u=f.get("u"),
                strike=f.get("strike", False), color=_color(f.get("color")), vertAlign=f.get("va"))
    fl = st.get("fill", {})
    fill = PatternFill(patternType=fl.get("pattern"), fgColor=_color(fl.get("fg")) or "00000000",
                       bgColor=_color(fl.get("bg")) or "00000000") if fl.get("pattern") else PatternFill()
    sides = {}
    for k in ("left", "right", "top", "bottom", "diagonal"):
        if k in st.get("border", {}):
            style, col = st["border"][k].split(" ", 1)
            sides[k] = Side(style=style, color=_color(col))
    diag = st.get("border", {}).get("diag", "")
    border = Border(diagonalUp="up" in diag, diagonalDown="down" in diag, **sides)
    a = st.get("align", {})
    align = Alignment(horizontal=a.get("h"), vertical=a.get("v"), wrap_text=a.get("wrap"), shrink_to_fit=a.get("shrink"),
                      indent=a.get("indent", 0), text_rotation=a.get("rot", 0))
    p = st.get("prot", {})
    prot = Protection(locked=not p.get("unlocked", False), hidden=p.get("hidden", False))
    return font, fill, border, align, prot, st.get("nf") or "General"


def cell_value(v):
    if isinstance(v, dict):
        if "f" in v:
            return v["f"]
        if "d" in v:
            try:
                return datetime.datetime.fromisoformat(v["d"])
            except ValueError:
                return v["d"]
    return v


def apply_overrides(wb, path):
    """목록을 openpyxl 통합문서에 적용. 반환: 적용한 항목 수. 파일이 없으면 0."""
    if not path or not os.path.exists(path):
        return 0
    with open(path, encoding="utf-8") as f:
        entries = json.load(f).get("entries", [])
    n = 0
    for e in entries:
        if e["sheet"] not in wb.sheetnames:
            continue
        ws = wb[e["sheet"]]
        if "merge" in e:
            ws.merge_cells(e["merge"])
        elif "unmerge" in e:
            if e["unmerge"] in {str(m) for m in ws.merged_cells.ranges}:
                ws.unmerge_cells(e["unmerge"])
        elif "col" in e:
            d = ws.column_dimensions[e["col"]]
            if "width" in e:
                d.width = e["width"]
            if "hidden" in e:
                d.hidden = bool(e["hidden"])
        elif "row" in e:
            d = ws.row_dimensions[int(e["row"])]
            if "height" in e:
                d.height = e["height"]
            if "hidden" in e:
                d.hidden = bool(e["hidden"])
        else:
            objs = style_objects(e["style"]) if "style" in e else None
            for rng in e.get("range", []):
                for r, c in expand(rng):
                    cell = ws.cell(row=r, column=c)
                    if "value" in e:
                        cell.value = cell_value(e["value"])
                    if objs:
                        cell.font, cell.fill, cell.border, cell.alignment, cell.protection, cell.number_format = objs
        n += 1
    return n
