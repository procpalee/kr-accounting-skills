# -*- coding: utf-8 -*-
"""xlsync_refs - 주소·수식 참조 도구.

- AxisMap: 기준본 행(열) 번호 -> 현재 번호. 엑셀이 행·열을 넣거나 지울 때 참조를 고치는 규칙과 같게 옮긴다
  (단일 칸이 지워지면 #REF!, 범위 끝점이 지워지면 안쪽으로 줄어듦, 범위 안에 넣은 줄은 범위에 포함).
- translate_formula: 수식 안의 A1 참조(다른 시트 포함)를 AxisMap 으로 옮긴 텍스트. 이름·구조적 참조는 그대로.
- group: 칸 목록 -> 'B5:B9' 같은 범위 문자열(행 우선으로 붙이고 같은 열 폭끼리 세로로 합침).
"""
import re
from bisect import bisect_left, bisect_right

from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter

CELL = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")
COLS = re.compile(r"^(\$?)([A-Za-z]{1,3}):(\$?)([A-Za-z]{1,3})$")
ROWS = re.compile(r"^(\$?)(\d+):(\$?)(\d+)$")
MAX_ROW, MAX_COL = 1048576, 16384


def rc(addr):
    """'B12' -> (12, 2)."""
    m = CELL.match(addr.replace("$", ""))
    return int(m.group(4)), column_index_from_string(m.group(2).upper())


def a1(r, c):
    return "%s%d" % (get_column_letter(c), r)


def split_sheet(ref):
    """"'시트 이름'!A1" -> ("시트 이름", "A1", "'시트 이름'!")."""
    if "!" not in ref:
        return None, ref, ""
    s, addr = ref.rsplit("!", 1)
    name = s[1:-1].replace("''", "'") if s.startswith("'") and s.endswith("'") else s
    return name, addr, s + "!"


def quote_sheet(name):
    return "'%s'" % name.replace("'", "''") if re.search(r"[^0-9A-Za-z_가-힣]", name) or name[:1].isdigit() else name


class AxisMap:
    """기준 번호 -> 현재 번호. keep = {기준: 현재}(남은 줄만), n = 기준 사용 범위 끝(그 안에서 keep 에 없으면 삭제된 줄).
    사용 범위 밖은 마지막 이동량만큼 민다."""

    def __init__(self, keep=None, n=None):
        self.keep = dict(keep or {})
        self.alive = sorted(self.keep)
        self.n = n if n is not None else (self.alive[-1] if self.alive else 0)
        self.identity = all(k == v for k, v in self.keep.items()) and len(self.keep) >= self.n

    def _outside(self, i):
        if not self.alive:
            return i
        if i > self.n:
            return i + self.keep[self.alive[-1]] - self.alive[-1]
        if i < self.alive[0] and i < 1:
            return i
        return None

    def one(self, i):
        if self.identity:
            return i
        if i in self.keep:
            return self.keep[i]
        return self._outside(i)

    def start(self, i):
        if self.identity or i in self.keep:
            return self.one(i)
        o = self._outside(i)
        if o is not None:
            return o
        j = bisect_left(self.alive, i)
        return self.keep[self.alive[j]] if j < len(self.alive) else None

    def end(self, i):
        if self.identity or i in self.keep:
            return self.one(i)
        o = self._outside(i)
        if o is not None:
            return o
        j = bisect_right(self.alive, i) - 1
        return self.keep[self.alive[j]] if j >= 0 else None


def translate_ref(text, cur_sheet, maps, renames=None):
    """참조 토큰 하나를 옮긴다. maps = {시트 이름(기준): (행 AxisMap, 열 AxisMap)}, renames = {기준 이름: 현재 이름}."""
    sheet, addr, prefix = split_sheet(text)
    target = sheet if sheet is not None else cur_sheet
    if sheet is not None and renames and sheet in renames:
        prefix = quote_sheet(renames[sheet]) + "!"
    m = maps.get(target)
    if m is None:
        return prefix + addr
    rm, cm = m
    parts = addr.split(":")
    if len(parts) == 1:
        g = CELL.match(addr)
        if not g:
            return prefix + addr
        r, c = rm.one(int(g.group(4))), cm.one(column_index_from_string(g.group(2).upper()))
        if r is None or c is None:
            return prefix + "#REF!"
        return prefix + "%s%s%s%d" % (g.group(1), get_column_letter(c), g.group(3), r)
    if len(parts) == 2:
        p, q = CELL.match(parts[0]), CELL.match(parts[1])
        if p and q:
            r1, r2 = rm.start(int(p.group(4))), rm.end(int(q.group(4)))
            c1 = cm.start(column_index_from_string(p.group(2).upper()))
            c2 = cm.end(column_index_from_string(q.group(2).upper()))
            if None in (r1, r2, c1, c2) or r1 > r2 or c1 > c2:
                return prefix + "#REF!"
            return prefix + "%s%s%s%d:%s%s%s%d" % (p.group(1), get_column_letter(c1), p.group(3), r1,
                                                     q.group(1), get_column_letter(c2), q.group(3), r2)
        g = COLS.match(addr)
        if g:
            c1 = cm.start(column_index_from_string(g.group(2).upper()))
            c2 = cm.end(column_index_from_string(g.group(4).upper()))
            if None in (c1, c2) or c1 > c2:
                return prefix + "#REF!"
            return prefix + "%s%s:%s%s" % (g.group(1), get_column_letter(c1), g.group(3), get_column_letter(c2))
        g = ROWS.match(addr)
        if g:
            r1, r2 = rm.start(int(g.group(2))), rm.end(int(g.group(4)))
            if None in (r1, r2) or r1 > r2:
                return prefix + "#REF!"
            return prefix + "%s%d:%s%d" % (g.group(1), r1, g.group(3), r2)
    return prefix + addr


def _tokens(formula):
    try:
        items = Tokenizer(formula).items
    except Exception:
        return None
    if "".join(t.value for t in items) != formula[1:]:        # 토큰을 이어 붙여 원문이 안 나오면 손대지 않는다
        return None
    return items


def translate_formula(formula, cur_sheet, maps, renames=None):
    if not isinstance(formula, str) or not formula.startswith("=") or not (maps or renames):
        return formula
    items = _tokens(formula)
    if items is None:
        return formula
    out = []
    for t in items:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            out.append(translate_ref(t.value, cur_sheet, maps, renames))
        else:
            out.append(t.value)
    return "=" + "".join(out)


def strip_refs(formula):
    """정렬용 서명: 칸 참조만 <r> 로 바꾼 수식(이름·문자열·함수는 남김)."""
    items = _tokens(formula) if isinstance(formula, str) and formula.startswith("=") else None
    if items is None:
        return formula
    out = []
    for t in items:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            _, addr, _ = split_sheet(t.value)
            out.append("<r>" if (CELL.match(addr) or COLS.match(addr) or ROWS.match(addr) or ":" in addr) else t.value)
        else:
            out.append(t.value)
    return "=" + "".join(out)


def literals(formula):
    """수식 안의 문자열 상수(따옴표 안) 목록."""
    items = _tokens(formula) if isinstance(formula, str) and formula.startswith("=") else None
    if items is None:
        return []
    return [t.value[1:-1].replace('""', '"') for t in items if t.type == Token.OPERAND and t.subtype == Token.TEXT and len(t.value) > 2]


def group(cells):
    """[(행, 열)] -> ['B5:B9', 'D2']."""
    rows = {}
    for r, c in sorted(set(cells)):
        rows.setdefault(r, []).append(c)
    segs = []
    for r in sorted(rows):
        cs = rows[r]
        s = p = cs[0]
        for c in cs[1:]:
            if c == p + 1:
                p = c
                continue
            segs.append((r, s, p))
            s = p = c
        segs.append((r, s, p))
    boxes, last = [], {}
    for r, c1, c2 in segs:
        k = (c1, c2)
        if k in last and boxes[last[k]][1] == r - 1:
            boxes[last[k]][1] = r
        else:
            boxes.append([r, r, c1, c2])
            last[k] = len(boxes) - 1
    out = []
    for r1, r2, c1, c2 in sorted(boxes, key=lambda b: (b[0], b[2])):
        out.append(a1(r1, c1) if (r1, c1) == (r2, c2) else "%s:%s" % (a1(r1, c1), a1(r2, c2)))
    return out


def expand(ref, cap=200000):
    """'A1:B3 D5' -> [(행, 열)...]. 너무 크면 None."""
    out = []
    for part in str(ref).split():
        part = part.replace("$", "")
        if ":" in part:
            p, q = part.split(":")
            (r1, c1), (r2, c2) = rc(p), rc(q)
        else:
            r1, c1 = r2, c2 = rc(part)
        if (r2 - r1 + 1) * (c2 - c1 + 1) + len(out) > cap:
            return None
        out.extend((r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1))
    return out
