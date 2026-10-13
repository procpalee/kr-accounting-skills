# -*- coding: utf-8 -*-
"""xlsync_snapshot - 엑셀 파일 -> 정규화한 스냅샷(dict, JSON 저장 가능). 읽기만 한다(원본을 저장하지 않음).

openpyxl: 값·수식·서식(글꼴·채우기·테두리·맞춤·보호·표시형식)·병합·데이터 유효성·조건부 서식·이름·메모·링크·시트 속성.
원본 XML: 열너비·사용자 지정 행높이(엑셀이 자동 맞춤으로 적은 높이는 제외)·숨김 행·셀 체크박스(엑셀 365)·그림/차트/도형 수.
이진 파트: 시트에 붙은 XML 아닌 파트(엑셀 COM 인쇄 설정이 만드는 printerSettings*.bin·OLE 개체 등)는 내용 대신 sha256 만.
olefile: VBA 모듈 소스(MS-OVBA 압축 해제). 계산 결과 캐시·선택 칸·확대율·수정 시각은 비교 대상이 아니다.
"""
import datetime
import hashlib
import io
import json
import os
import re
import struct
import zipfile
from html import unescape

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

from xlsync_refs import a1, expand, get_column_letter

VERSION = 2          # 2: 시트마다 이진 파트 해시 "bin"(2026-09-30). 1 기준본에는 없어 비교를 건너뛴다


def color(c):
    if c is None:
        return None
    t = getattr(c, "type", None)
    if t == "rgb":
        v = c.rgb if isinstance(c.rgb, str) else None
        if not v:
            return None
        v = v.upper()
        return v[2:] if len(v) == 8 else v
    if t == "theme":
        return "theme%d%s" % (c.theme, ("%+.3f" % c.tint) if c.tint else "")
    if t == "indexed":
        return None if c.indexed in (64, 65) else "idx%d" % c.indexed
    return None


def _clean(d):
    return {k: v for k, v in d.items() if v not in (None, False, 0, "", [], {})}


def style_of(cell):
    f, fl, b, al, p = cell.font, cell.fill, cell.border, cell.alignment, cell.protection
    font = _clean({"name": f.name, "sz": float(f.sz) if f.sz else None, "b": bool(f.b), "i": bool(f.i),
                   "u": f.u, "strike": bool(f.strike), "color": color(f.color), "va": f.vertAlign}) if f else {}
    fill = {}
    if fl is not None and fl.__class__.__name__ == "GradientFill":
        fill = {"gradient": "%s:%s" % (fl.type, [(s.position, color(s.color)) for s in fl.stop])}
    elif fl is not None and fl.patternType:
        fill = {"pattern": fl.patternType, "fg": color(fl.fgColor)}
        if fl.patternType != "solid":
            fill["bg"] = color(fl.bgColor)
    border = {}
    if b is not None:
        for k in ("left", "right", "top", "bottom", "diagonal"):
            s = getattr(b, k)
            if s is not None and s.style:
                border[k] = "%s %s" % (s.style, color(s.color) or "auto")
        if "diagonal" in border:
            border["diag"] = "%s%s" % ("up" if b.diagonalUp else "", "down" if b.diagonalDown else "")
    align = _clean({"h": None if (al is None or al.horizontal == "general") else al.horizontal,
                    "v": None if (al is None or al.vertical in (None, "bottom")) else al.vertical,
                    "wrap": bool(al and al.wrap_text), "shrink": bool(al and al.shrink_to_fit),
                    "indent": int(al.indent or 0) if al else 0, "rot": int(al.text_rotation or 0) if al else 0})
    prot = _clean({"unlocked": bool(p is not None and p.locked is False), "hidden": bool(p is not None and p.hidden)})
    return _clean({"font": font, "fill": fill, "border": border, "align": align, "prot": prot,
                   "nf": nf_norm(cell.number_format)})


def nf_norm(nf):
    """표시형식 정규화: 엑셀은 저장할 때 날짜 구분자 앞에 역슬래시를 붙인다(yyyy-mm-dd 가 yyyy역슬래시-mm... 로 저장됨)."""
    if not nf or nf == "General":
        return None
    return re.sub(r"\\([^\\])", r"\1", nf)


def sid(style, table):
    key = hashlib.sha1(json.dumps(style, sort_keys=True, ensure_ascii=False).encode("utf-8")).hexdigest()[:10]
    table.setdefault(key, style)
    return key


def num(v):
    if isinstance(v, float):
        if v.is_integer() and abs(v) < 1e15:
            return int(v)
        return float("%.15g" % v)
    return v


def value_of(cell):
    v = cell.value
    if v is None:
        return None
    name = v.__class__.__name__
    if name == "ArrayFormula":
        return {"f": "{%s}" % (v.text or ""), "ref": str(v.ref)}
    if name == "DataTableFormula":
        return {"f": "{TABLE}"}
    if cell.data_type == "f" and isinstance(v, str):
        return {"f": re.sub(r"_xl(?:fn|ws)\.", "", v)}
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return num(v)
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return {"d": v.isoformat()}
    if isinstance(v, datetime.timedelta):
        return {"d": str(v)}
    return str(v).replace("\r\n", "\n")


def _attrs(tag):
    return {k: unescape(v) for k, v in re.findall(r'([\w:]+)="([^"]*)"', tag)}


def _read(z, name):
    """XML 파트를 글자로(없으면 ""). UTF-8 이 아닌 바이트는 대체 문자로 - 이진 파트가 섞여도 멈추지 않는다."""
    try:
        return z.read(name).decode("utf-8", "replace")
    except KeyError:
        return ""


def _bin_sig(data):
    """이진 파트는 풀어 읽지 않고 sha256 앞 16자로만 같음/바뀜을 본다."""
    return "sha256:" + hashlib.sha256(data).hexdigest()[:16]


def raw_parts(path):
    """시트 이름 -> 원본 XML 에서 읽는 항목(열·사용자 행높이·체크박스·도형 수·이진 파트 해시)."""
    out = {}
    with zipfile.ZipFile(path) as z:
        parts = set(z.namelist())
        rels = {a.get("Id"): a.get("Target") for a in map(_attrs, re.findall(r"<Relationship\b[^>]*>", _read(z, "xl/_rels/workbook.xml.rels")))}
        styles = _read(z, "xl/styles.xml")
        cbx = set()
        if 'type="Checkbox"' in _read(z, "xl/featurePropertyBag/featurePropertyBag.xml"):
            block = re.search(r"<cellXfs\b[^>]*>(.*?)</cellXfs>", styles, re.S)
            xfs = re.findall(r"<xf\b[^>]*/>|<xf\b[^>]*>.*?</xf>", block.group(1), re.S) if block else []
            cbx = {i for i, x in enumerate(xfs) if "xfComplement" in x}
        for tag in re.findall(r"<sheet\b[^>]*>", _read(z, "xl/workbook.xml")):
            a = _attrs(tag)
            target = rels.get(a.get("r:id"), "")
            part = target.lstrip("/") if target.startswith("/") else "xl/" + target
            xml = _read(z, part)
            info = {"cols": {}, "rows": {}, "checkbox": {}, "drawing": {}, "bin": {}}
            for c in map(_attrs, re.findall(r"<col\b[^>]*>", xml)):
                for i in range(int(c.get("min", 0)), min(int(c.get("max", 0)), 16384) + 1):
                    d = {}
                    if c.get("width") and c.get("customWidth") in ("1", "true"):
                        d["w"] = round(float(c["width"]), 4)
                    if c.get("hidden") in ("1", "true"):
                        d["hidden"] = True
                    if d:
                        info["cols"][get_column_letter(i)] = d
            for r in map(_attrs, re.findall(r"<row\b[^>]*>", xml)):
                d = {}
                if r.get("ht") and r.get("customHeight") in ("1", "true"):
                    d["h"] = round(float(r["ht"]), 2)
                if r.get("hidden") in ("1", "true"):
                    d["hidden"] = True
                if d:
                    info["rows"][r.get("r")] = d
            if cbx:
                for m in re.finditer(r'<c\b([^>]*?)(/>|>(.*?)</c>)', xml, re.S):
                    a2 = _attrs(m.group(1))
                    if a2.get("s") and int(a2["s"]) in cbx:
                        v = re.search(r"<v>([^<]*)</v>", m.group(3) or "")
                        info["checkbox"][a2.get("r")] = bool(v and v.group(1) in ("1", "TRUE", "true"))
            srels = _read(z, re.sub(r"([^/]+)$", r"_rels/\1.rels", part))
            for t in map(_attrs, re.findall(r"<Relationship\b[^>]*>", srels)):
                tp, tg = t.get("Type", ""), t.get("Target", "")
                dpart = os.path.normpath(os.path.join(os.path.dirname(part), tg)).replace("\\", "/")
                if tp.endswith("/drawing"):
                    dx = _read(z, dpart)
                    info["drawing"]["pic"] = dx.count("<xdr:pic>")
                    info["drawing"]["chart"] = dx.count("drawingml/2006/chart")
                    info["drawing"]["shape"] = dx.count("<xdr:sp>") + dx.count("<xdr:sp ")
                elif tp.endswith("/vmlDrawing"):
                    info["drawing"]["vml"] = _read(z, dpart).count("<v:shape ")
                elif dpart in parts and not dpart.lower().endswith((".xml", ".vml")):     # printerSettings*.bin·oleObject*.bin·배경 그림 등
                    info["bin"].setdefault(tp.rsplit("/", 1)[-1], []).append(_bin_sig(z.read(dpart)))
            info["drawing"] = {k: v for k, v in info["drawing"].items() if v}
            info["bin"] = {k: ", ".join(sorted(v)) for k, v in info["bin"].items()}      # 파트 번호는 저장 때마다 바뀔 수 있어 관계 종류로 묶음
            out[a.get("name")] = info
    return out


# ── VBA(MS-OVBA) ──────────────────────────────────────────────────────────
def ovba_decompress(buf):
    if not buf or buf[0] != 1:
        raise ValueError("VBA 압축 컨테이너 아님")
    out, pos, n = bytearray(), 1, len(buf)
    while pos + 2 <= n:
        hdr = struct.unpack_from("<H", buf, pos)[0]
        end = min(pos + (hdr & 0x0FFF) + 3, n)
        pos += 2
        start = len(out)
        if not (hdr >> 15) & 1:
            out += buf[pos:pos + 4096]
            pos += 4096
            continue
        while pos < end:
            flags = buf[pos]
            pos += 1
            for bit in range(8):
                if pos >= end:
                    break
                if not (flags >> bit) & 1:
                    out.append(buf[pos])
                    pos += 1
                    continue
                tok = struct.unpack_from("<H", buf, pos)[0]
                pos += 2
                bc = max((len(out) - start - 1).bit_length(), 4)
                length, offset = (tok & (0xFFFF >> bc)) + 3, (tok >> (16 - bc)) + 1
                for _ in range(length):
                    out.append(out[-offset])
    return bytes(out)


def vba_modules(path):
    with zipfile.ZipFile(path) as z:
        if "xl/vbaProject.bin" not in z.namelist():
            return None
        data = z.read("xl/vbaProject.bin")
    try:
        import olefile
        ole = olefile.OleFileIO(io.BytesIO(data))
        d = ovba_decompress(ole.openstream("VBA/dir").read())
        i, cp, mods, cur = 0, 1252, [], {}
        while i + 6 <= len(d):
            rid, size = struct.unpack_from("<HI", d, i)
            i += 6
            if rid == 0x0009:
                size = 6
            rec = d[i:i + size]
            i += size
            if rid == 0x0003:
                cp = struct.unpack_from("<H", rec)[0]
            elif rid == 0x0019:
                cur = {"name": rec}
            elif rid == 0x0047:
                cur["uname"] = rec
            elif rid == 0x001A:
                cur["stream"] = rec
            elif rid == 0x0031:
                cur["offset"] = struct.unpack_from("<I", rec)[0]
            elif rid == 0x002B and cur:
                mods.append(cur)
                cur = {}
            elif rid == 0x0010:
                break
        enc = "cp%d" % cp
        out = {}
        for m in mods:
            name = m["uname"].decode("utf-16-le") if m.get("uname") else m["name"].decode(enc, "replace")
            src = ovba_decompress(ole.openstream("VBA/" + m["stream"].decode(enc, "replace")).read()[m.get("offset", 0):])
            lines = [ln for ln in src.decode(enc, "replace").replace("\r\n", "\n").split("\n") if not ln.startswith("Attribute VB_")]
            out[name] = "\n".join(lines).strip("\n")
        return out
    except Exception as e:  # 읽지 못하면 이진 해시로만 비교
        return {"__bin__": "읽기 실패(%s) sha1 %s" % (e.__class__.__name__, hashlib.sha1(data).hexdigest()[:12])}


# ── 스냅샷 ────────────────────────────────────────────────────────────────
def _dv_sig(dv):
    def f(x):
        return x[1:] if isinstance(x, str) and x.startswith("=") else x      # 엑셀은 목록 원본의 '=' 를 빼고 저장
    return _clean({"type": dv.type, "op": dv.operator, "f1": f(dv.formula1), "f2": f(dv.formula2),
                   "blank": bool(dv.allow_blank), "err": bool(dv.showErrorMessage), "inp": bool(dv.showInputMessage),
                   "errStyle": dv.errorStyle, "errT": dv.errorTitle, "err_msg": dv.error, "prT": dv.promptTitle, "pr": dv.prompt})


def snapshot(path, label=""):
    raw = raw_parts(path)
    wb = load_workbook(path, data_only=False, keep_vba=False, rich_text=False, keep_links=False)
    styles, dvs = {}, {}
    st = os.stat(path)
    snap = {"meta": {"xlsync": VERSION, "file": os.path.abspath(path), "label": label,
                     "mtime": datetime.datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"),
                     "sha1": hashlib.sha1(open(path, "rb").read()).hexdigest(),
                     "made": datetime.datetime.now().isoformat(timespec="seconds")},
            "sheets": [], "names": {}, "vba": vba_modules(path), "styles": styles, "dvs": dvs}
    for n, dn in wb.defined_names.items():
        snap["names"][n] = dn.attr_text
    for ws in wb.worksheets:
        for n, dn in ws.defined_names.items():
            snap["names"]["%s!%s" % (ws.title, n)] = dn.attr_text
        default = sid(style_of(Cell(ws)), styles)
        r = raw.get(ws.title, {})
        sh = {"name": ws.title, "state": ws.sheet_state, "tab": color(ws.sheet_properties.tabColor),
              "freeze": ws.freeze_panes, "grid": ws.sheet_view.showGridLines is not False,
              "print": _clean({"area": ws.print_area if ws.print_area else None, "rows": ws.print_title_rows, "cols": ws.print_title_cols}),
              "default": default, "cols": r.get("cols", {}), "rows": r.get("rows", {}),
              "checkbox": r.get("checkbox", {}), "drawing": r.get("drawing", {}), "bin": r.get("bin", {}),
              "merges": sorted(str(m) for m in ws.merged_cells.ranges), "cells": {}, "dv": {}, "dv_big": [], "cf": [],
              "comments": {}, "links": {}}
        for (row, col), c in ws._cells.items():
            v, s = value_of(c), sid(style_of(c), styles)
            if v is None and s == default:
                continue
            e = {"s": s}
            if v is not None:
                e["v"] = v
            sh["cells"][a1(row, col)] = e
            if getattr(c, "comment", None) is not None:
                sh["comments"][a1(row, col)] = c.comment.text
            if getattr(c, "hyperlink", None) is not None:
                h = c.hyperlink
                sh["links"][a1(row, col)] = h.target or ("#" + (h.location or ""))
        for dv in ws.data_validations.dataValidation:
            key = sid(_dv_sig(dv), dvs)
            cells = expand(dv.sqref)
            if cells is None:
                sh["dv_big"].append([str(dv.sqref), key])
                continue
            for rr, cc in cells:
                sh["dv"][a1(rr, cc)] = key
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                dxf = {}
                if rule.dxf is not None:
                    dxf = _clean({"font": color(rule.dxf.font.color) if rule.dxf.font else None,
                                  "b": bool(rule.dxf.font and rule.dxf.font.b),
                                  "fill": color(rule.dxf.fill.fgColor) if rule.dxf.fill is not None and rule.dxf.fill.patternType else None,
                                  "fill_bg": color(rule.dxf.fill.bgColor) if rule.dxf.fill is not None else None,
                                  "nf": rule.dxf.numFmt.formatCode if rule.dxf.numFmt else None})
                sig = _clean({"type": rule.type, "op": rule.operator, "f": list(rule.formula or []), "text": rule.text,
                              "stop": bool(rule.stopIfTrue), "dxf": dxf})
                sh["cf"].append([str(cf.sqref), sig])
        sh["cf"].sort(key=lambda x: json.dumps(x, sort_keys=True, ensure_ascii=False))
        snap["sheets"].append(sh)
    for cs in wb.chartsheets:
        snap["sheets"].append({"name": cs.title, "chartsheet": True, "state": cs.sheet_state})
    wb.close()
    return snap


def load_or_snap(path):
    """.json 이면 스냅샷 파일, 엑셀이면 새로 읽는다."""
    if path.lower().endswith(".json"):
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    return snapshot(path)
