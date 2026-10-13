# -*- coding: utf-8 -*-
"""xlsync_trace - 생성 스크립트를 실행하면서 openpyxl 이 칸·열·행·병합·유효성·조건부서식·이름을 쓸 때마다
'그 값을 쓴 코드 줄'(파일:줄 함수 <- 호출한 곳)을 기록한다. 저장한 통합문서마다 기록 파일(prov)을 남긴다.

    python -X utf8 xlsync.py record --prov out.prov.json -- build.py [인자...]

기록은 실행 흐름 그대로라 결정론적이다. openpyxl 로 다른 파일을 읽는 동안(load_workbook)은 기록하지 않는다.
"""
import datetime
import json
import os
import runpy
import sys
import sysconfig

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
_SKIP_DIRS = tuple(os.path.normcase(p) + os.sep for p in {
    os.path.dirname(openpyxl.__file__), os.path.dirname(os.__file__), sysconfig.get_paths().get("purelib", ""),
    sysconfig.get_paths().get("platlib", ""), HERE} if p)


class Recorder:
    def __init__(self, root):
        self.root = root
        self.active = True
        self.sites, self.site_ix = [], {}
        self.cells, self.dims, self.merges, self.titles = {}, {}, {}, {}
        self.dv, self.dv_cells, self.cf, self.names = {}, {}, {}, {}
        self.ws = {}
        self.files = {}
        self._skip = {}

    def _skipped(self, fn):
        r = self._skip.get(fn)
        if r is None:
            n = os.path.normcase(os.path.abspath(fn)) if not fn.startswith("<") else fn
            r = fn.startswith("<") or n.startswith(_SKIP_DIRS)
            self._skip[fn] = r
        return r

    def _rel(self, fn):
        a = os.path.abspath(fn)
        try:
            rel = os.path.relpath(a, self.root)
            if not rel.startswith(".."):
                return rel.replace("\\", "/")
        except ValueError:
            pass
        home = os.path.expanduser("~")
        return ("~/" + os.path.relpath(a, home)).replace("\\", "/") if a.startswith(home) else a.replace("\\", "/")

    def site(self, depth=2):
        f = sys._getframe(depth)
        chain = []
        while f is not None and len(chain) < 3:
            fn = f.f_code.co_filename
            if not self._skipped(fn):
                chain.append("%s:%d %s" % (self._rel(fn), f.f_lineno, f.f_code.co_name))
            f = f.f_back
        s = " <- ".join(chain) or "?"
        ix = self.site_ix.get(s)
        if ix is None:
            ix = self.site_ix[s] = len(self.sites)
            self.sites.append(s)
        return ix


REC = None
_CELL_ATTRS = {"value": "value", "font": "font", "fill": "fill", "border": "border", "alignment": "alignment",
               "number_format": "number_format", "protection": "protection", "style": "style", "_style": "style",
               "comment": "comment", "hyperlink": "hyperlink"}


def install(root):
    """openpyxl 에 기록 고리를 건다(한 번만)."""
    global REC
    if REC is not None:
        return REC
    REC = rec = Recorder(root)
    from openpyxl.cell.cell import Cell, MergedCell
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.reader import excel as reader
    from openpyxl.workbook.child import _WorkbookChild
    from openpyxl.workbook.defined_name import DefinedNameDict
    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.dimensions import ColumnDimension, RowDimension
    from openpyxl.worksheet.worksheet import Worksheet

    def cell_setattr(self, name, value):
        object.__setattr__(self, name, value)
        a = _CELL_ATTRS.get(name)
        if a is None or not rec.active:
            return
        if sys._getframe(1).f_code.co_name == "__init__" and name == "_style":
            return
        try:
            ws = self.parent
            rec.ws[id(ws)] = ws
            rec.cells.setdefault(id(ws), {}).setdefault(self.coordinate, {})[a] = rec.site()
        except Exception:
            pass
    Cell.__setattr__ = cell_setattr
    MergedCell.__setattr__ = cell_setattr

    def dim_setattr(cls, kind, names):
        base = cls.__setattr__

        def f(self, name, value):
            base(self, name, value)
            if name in names and rec.active and sys._getframe(1).f_code.co_name != "__init__":
                ws, idx = getattr(self, "parent", None), getattr(self, "index", None)
                if ws is not None and idx is not None:
                    rec.ws[id(ws)] = ws
                    rec.dims.setdefault(id(ws), {}).setdefault((kind, str(idx)), {})[names[name]] = rec.site()
        cls.__setattr__ = f
    dim_setattr(ColumnDimension, "col", {"width": "width", "hidden": "hidden"})
    dim_setattr(RowDimension, "row", {"ht": "height", "height": "height", "hidden": "hidden"})

    orig_merge = Worksheet.merge_cells

    def merge_cells(self, range_string=None, start_row=None, start_column=None, end_row=None, end_column=None):
        out = orig_merge(self, range_string, start_row, start_column, end_row, end_column)
        if rec.active:
            from openpyxl.worksheet.cell_range import CellRange
            rng = str(CellRange(range_string, min_col=start_column, min_row=start_row, max_col=end_column, max_row=end_row)) if range_string is None else str(range_string).upper()
            rec.ws[id(self)] = self
            rec.merges.setdefault(id(self), {})[rng] = rec.site()
        return out
    Worksheet.merge_cells = merge_cells

    title = _WorkbookChild.title

    def set_title(self, value):
        title.fset(self, value)
        if rec.active:
            rec.titles[id(self)] = rec.site()
    _WorkbookChild.title = property(title.fget, set_title)

    orig_dv_add = DataValidation.add

    def dv_add(self, cell):
        out = orig_dv_add(self, cell)
        if rec.active:
            rec.dv_cells.setdefault(id(self), []).append([str(getattr(cell, "coordinate", cell)), rec.site()])
        return out
    DataValidation.add = dv_add
    orig_add_dv = Worksheet.add_data_validation

    def add_data_validation(self, dv):
        out = orig_add_dv(self, dv)
        if rec.active:
            rec.dv[id(dv)] = rec.site()
        return out
    Worksheet.add_data_validation = add_data_validation

    orig_cf_add = ConditionalFormattingList.add

    def cf_add(self, range_string, cfRule):
        out = orig_cf_add(self, range_string, cfRule)
        if rec.active:
            rec.cf.setdefault(id(self), []).append([str(range_string), rec.site()])
        return out
    ConditionalFormattingList.add = cf_add

    orig_set = DefinedNameDict.__setitem__

    def name_set(self, key, value):
        orig_set(self, key, value)
        if rec.active:
            rec.names[str(key)] = rec.site()
    DefinedNameDict.__setitem__ = name_set

    orig_load = reader.load_workbook

    def load_workbook(*a, **kw):
        prev, rec.active = rec.active, False
        try:
            return orig_load(*a, **kw)
        finally:
            rec.active = prev
    reader.load_workbook = load_workbook
    openpyxl.load_workbook = load_workbook

    orig_save = Workbook.save

    def save(self, filename):
        prev, rec.active = rec.active, False
        try:
            orig_save(self, filename)
        finally:
            rec.active = prev
        rec.files[os.path.abspath(str(filename))] = dump(self)
    Workbook.save = save
    return rec


def dump(wb):
    rec = REC
    out = {"sheets": {}, "names": {n: rec.names[n] for n in wb.defined_names.keys() if n in rec.names}}
    for ws in wb.worksheets:
        k = id(ws)
        sh = {"cells": rec.cells.get(k, {}), "cols": {}, "rows": {}, "merges": rec.merges.get(k, {}),
              "dv": [], "cf": [], "title": rec.titles.get(k)}
        for (kind, idx), attrs in rec.dims.get(k, {}).items():
            sh["cols" if kind == "col" else "rows"][idx] = attrs
        for dv in ws.data_validations.dataValidation:
            if id(dv) in rec.dv or id(dv) in rec.dv_cells:
                sh["dv"].append({"sqref": str(dv.sqref), "made": rec.dv.get(id(dv)), "cells": rec.dv_cells.get(id(dv), [])})
        sh["cf"] = rec.cf.get(id(ws.conditional_formatting), [])
        out["sheets"][ws.title] = sh
    return out


def run(prov_path, argv):
    script = os.path.abspath(argv[0])
    root = os.path.dirname(script)
    rec = install(root)
    sys.argv = [script] + list(argv[1:])
    sys.path.insert(0, root)
    code = 0
    try:
        runpy.run_path(script, run_name="__main__")
    except SystemExit as e:
        code = e.code if isinstance(e.code, int) else (0 if e.code is None else 1)
    data = {"xlsync_prov": 1, "script": script, "root": root, "argv": argv[1:],
            "made": datetime.datetime.now().isoformat(timespec="seconds"), "sites": rec.sites, "files": rec.files}
    with open(prov_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)
    return code, data


def load(prov_path, target=None):
    """기록 파일에서 target(최종 파일)에 해당하는 통합문서 기록. 이름(확장자 제외)이 같은 것, 없으면 마지막 저장본."""
    with open(prov_path, encoding="utf-8") as f:
        data = json.load(f)
    files = data.get("files", {})
    if not files:
        return None
    pick = None
    if target:
        stem = os.path.splitext(os.path.basename(target))[0].lower()
        for p in files:
            if os.path.splitext(os.path.basename(p))[0].lower() == stem:
                pick = p
    pick = pick or list(files)[-1]
    entry = dict(files[pick])
    entry["sites"], entry["file"], entry["script"], entry["root"] = data["sites"], pick, data.get("script"), data.get("root")
    return entry
