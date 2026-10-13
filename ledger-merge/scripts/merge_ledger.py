"""더존 계정별원장 엑셀(계정마다 시트 1개)을 시트 하나로 합친다.

사용: python merge_ledger.py <원장.xlsx> [--out-dir 03_산출 경로] [--overwrite]
결과: <03_산출>/<원래 파일 이름>_통합.xlsx  (시트: 통합 · 확인)
- [월계]·[누계] 줄은 뺀다. 전기이월 줄은 남긴다(일자 없음 → 빈칸).
- 확인 시트: 계정별 통합 차변·대변 합계(수식)를 원장 마지막 [누계] 줄과 비교.
"""
import argparse
import datetime as dt
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl import Workbook

sys.path.insert(0, str(Path.home() / ".claude" / "skills" / "excel-theme"))
from excel_theme import apply_theme, get_theme, style_total_row  # noqa: E402

THEME = "default"
HEADERS = ["계정코드", "계정명", "일자", "적요", "거래처코드", "거래처", "차변", "대변", "잔액"]
CHECK_HEADERS = ["계정코드", "계정명", "통합 차변", "통합 대변", "원장 누계 차변", "원장 누계 대변",
                 "차변 차이", "대변 차이", "결과"]
SRC_HEAD = ["날짜", "적요란", "코드", "거래처", "차변", "대변", "잔액"]  # 원장 열 순서(공백 제거 후 비교)
HEAD_ROW, FIRST_ROW = 4, 5


def squash(v):
    return re.sub(r"\s+", "", str(v)) if v is not None else ""


def find_account(ws):
    """시트 상단의 '계정과목 : [코드] 이름' 에서 (코드, 이름). 없으면 시트명 '이름(코드)'."""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 6), values_only=True):
        for v in row:
            if v and "계정과목" in str(v):
                m = re.search(r"\[([^\]]+)\]\s*(.+)$", str(v))
                if m:
                    return m.group(1).strip(), m.group(2).strip()
    m = re.match(r"(.+)\((\w+)\)$", ws.title)
    return (m.group(2), m.group(1)) if m else (None, None)


def find_start_year(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 6), values_only=True):
            for v in row:
                m = re.search(r"(\d{4})\.\d{2}\.\d{2}\s*~", str(v)) if v else None
                if m:
                    return int(m.group(1))
    return None


def to_date(v, year, prev_month):
    """'MM-DD' 또는 날짜값 → (date, 월, 연). 월이 줄어들면 해가 바뀐 것으로 본다."""
    if isinstance(v, dt.datetime):
        return v.date(), v.month, v.year
    s = str(v).strip()
    m = re.match(r"^(\d{4})[-./](\d{2})[-./](\d{2})$", s)
    if m:
        y, mo, d = map(int, m.groups())
        return dt.date(y, mo, d), mo, y
    m = re.match(r"^(\d{2})[-./](\d{2})$", s)
    if not m:
        return None, prev_month, year
    mo, d = map(int, m.groups())
    if prev_month and mo < prev_month:
        year += 1
    return dt.date(year, mo, d), mo, year


def read_ledger(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    year0 = find_start_year(wb)
    rows, checks, warns = [], [], []
    for ws in wb.worksheets:
        code, name = find_account(ws)
        if code is None:
            warns.append(f"[{ws.title}] 계정과목을 찾지 못해 건너뜀 - 확인 필요")
            continue
        hdr = next((r for r in range(1, min(ws.max_row, 8) + 1) if squash(ws.cell(r, 1).value) == "날짜"), None)
        if hdr is None or [squash(ws.cell(hdr, c).value) for c in range(1, 8)] != SRC_HEAD:
            warns.append(f"[{ws.title}] 머리글이 예상({'/'.join(SRC_HEAD)})과 달라 건너뜀 - 확인 필요")
            continue
        year, prev_m, last_total = year0, None, None
        for r in ws.iter_rows(min_row=hdr + 1, max_col=7, values_only=True):
            if all(v is None for v in r):
                continue
            t = squash(r[1])
            if t in ("[월계]", "[누계]"):
                if t == "[누계]":
                    last_total = (r[4] or 0, r[5] or 0)
                continue
            d = None
            if r[0] is not None and year:
                d, prev_m, year = to_date(r[0], year, prev_m)
                if d is None:
                    warns.append(f"[{ws.title}] 일자 해석 불가 '{r[0]}' - 확인 필요")
            rows.append([code, name, d, r[1], r[2], r[3], r[4], r[5], r[6]])
        if last_total is None:
            warns.append(f"[{ws.title}] [누계] 줄이 없어 대조 불가 - 확인 필요")
        checks.append((code, name, last_total))
    return rows, checks, warns


def locate_out_dir(src, given):
    if given:
        return Path(given)
    for p in src.resolve().parents:
        if (p / "03_산출").is_dir():
            return p / "03_산출"
    sys.exit("03_산출 폴더를 찾지 못했다. --out-dir 로 지정할 것.")


def build(rows, checks, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "통합"
    ws["B2"] = "통합"
    for i, h in enumerate(HEADERS):
        ws.cell(HEAD_ROW, 2 + i, h)
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row):
            if v is not None:
                ws.cell(FIRST_ROW + ri, 2 + ci, v)
    last = FIRST_ROW + len(rows) - 1
    apply_theme(ws, THEME, header_row=HEAD_ROW, data_range=f"B{HEAD_ROW}:J{last}", title_cell="B2",
                number_format_cols={"D": "yyyy-mm-dd", "H": "accounting", "I": "accounting", "J": "accounting"},
                freeze=True)

    cs = wb.create_sheet("확인")
    cs["B2"] = "확인"
    for i, h in enumerate(CHECK_HEADERS):
        cs.cell(HEAD_ROW, 2 + i, h)

    def rng(col):
        return f"통합!${col}${FIRST_ROW}:${col}${last}"

    for i, (code, name, tot) in enumerate(checks):
        r = FIRST_ROW + i
        cs.cell(r, 2, code)
        cs.cell(r, 3, name)
        cs.cell(r, 4, f"=SUMIFS({rng('H')},{rng('B')},B{r})")
        cs.cell(r, 5, f"=SUMIFS({rng('I')},{rng('B')},B{r})")
        if tot is None:
            cs.cell(r, 6, "확인 필요")
            cs.cell(r, 7, "확인 필요")
        else:
            cs.cell(r, 6, tot[0])
            cs.cell(r, 7, tot[1])
        cs.cell(r, 8, f'=IF(ISNUMBER(F{r}),D{r}-F{r},"")')
        cs.cell(r, 9, f'=IF(ISNUMBER(G{r}),E{r}-G{r},"")')
        cs.cell(r, 10, f'=IF(COUNT(F{r}:G{r})<2,"확인 필요",IF(AND(H{r}=0,I{r}=0),"일치","불일치"))')
    n_last = FIRST_ROW + len(checks) - 1
    tr = n_last + 1
    cs.cell(tr, 2, "합계")
    for col in "DEFGHI":
        cs[f"{col}{tr}"] = f"=SUM({col}{FIRST_ROW}:{col}{n_last})"
    cs[f"J{tr}"] = (f'=IF(COUNTIF(J{FIRST_ROW}:J{n_last},"일치")=ROWS(J{FIRST_ROW}:J{n_last}),"전체 일치",'
                    f'COUNTIF(J{FIRST_ROW}:J{n_last},"<>일치")&"건 확인")')
    fmt = {c: "accounting" for c in "DEFGHI"}
    apply_theme(cs, THEME, header_row=HEAD_ROW, data_range=f"B{HEAD_ROW}:J{tr}", title_cell="B2",
                number_format_cols=fmt, freeze=True)
    style_total_row(cs, get_theme(THEME), row=tr, min_col=2, max_col=10)
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("--out-dir")
    ap.add_argument("--overwrite", action="store_true")
    a = ap.parse_args()
    src = Path(a.src)
    out = locate_out_dir(src, a.out_dir) / f"{src.stem}_통합.xlsx"
    if out.exists() and not a.overwrite:
        sys.exit(f"이미 있음: {out}\n덮어쓰려면 --overwrite (사용자가 손댄 파일인지 먼저 확인).")
    rows, checks, warns = read_ledger(src)
    if not rows:
        for w in warns[:5]:
            print("주의:", w)
        sys.exit("합칠 거래 줄이 없다 - 원장 형식 확인 필요")
    out.parent.mkdir(parents=True, exist_ok=True)
    build(rows, checks, out)

    # 파이썬 쪽 대조(보고용). 엑셀 확인 시트 수식과 같은 계산.
    sd, sc = {}, {}
    for r in rows:
        sd[r[0]] = sd.get(r[0], 0) + (r[6] or 0)
        sc[r[0]] = sc.get(r[0], 0) + (r[7] or 0)
    bad = [c for c, _, t in checks if t is not None and (sd.get(c, 0), sc.get(c, 0)) != tuple(t)]
    nochk = [c for c, _, t in checks if t is None]
    print(f"저장: {out}")
    print(f"계정 {len(checks)}개 · 거래 줄 {len(rows)}개")
    print(f"누계 대조: 일치 {len(checks) - len(bad) - len(nochk)} / 불일치 {len(bad)} / 확인 필요 {len(nochk)}")
    if bad:
        print("불일치 계정:", ", ".join(bad))
    for w in warns:
        print("주의:", w)


if __name__ == "__main__":
    main()
