"""워드(.docx) · 파워포인트(.pptx) 두 파일 비교 - 보고서 · 의견서처럼 문서가 결과물인 업무의 '맞추기'.

워드 : 문단 글자(순서대로) · 표 칸
PPT  : 쪽마다 도형 글자 · 표 칸 · 차트 값(항목 · 계열)
--numbers 를 주면 글자 대신 '글 속 숫자 모음'만 비교한다(문장 · 순서는 달라도 같은 금액이 같은 횟수만큼 나오는지).
글로 쓴 칸은 결론이 같으면 된다 - 문장 차이는 사람이 읽고 판단하고, 숫자 · 표 · 차트 값은 0 이 목표다.

사용: python -X utf8 compare_docs.py <만든 파일> <완성본> [--numbers] [--max 40] [--json 결과.json]
종료 코드: 차이 없음 0 · 차이 있음 1 · 실행 오류 2
"""
from __future__ import annotations

import argparse
import difflib
import json
import re
import sys
from collections import Counter
from pathlib import Path

# 숫자 앞이 숫자 · 소수점 · 쉼표가 아니면 숫자의 시작('매출액1,234' 의 1,234 도 잡는다). '-' 는 앞이 숫자가 아닐 때만 음수 표시.
NUM_RE = re.compile(r"(?<![\d.,])(?:(?<![\d])[-△])?\d[\d,]*(?:\.\d+)?%?")
UNIT_AFTER = re.compile(r"\s?(조|억|만|천|원|%|배|개월|년|건|명|개)")


def numbers_in(text: str) -> Counter:
    """글 속 숫자 모음(몇 번 나오는지). 쉼표는 지운다. 한 자리 숫자는 단위가 붙었을 때만 센다('7억'은 세고 '① · 3쪽'은 뺀다)."""
    bag = Counter()
    text = text or ""
    for m in NUM_RE.finditer(text):
        t = m.group(0).replace(",", "").replace("△", "-")
        if re.fullmatch(r"-?\d", t) and not UNIT_AFTER.match(text, m.end()):
            continue
        bag[t] += 1
    return bag


def diff_bags(a: Counter, b: Counter, where: str, out: list):
    """숫자 모음 비교 - 순서가 달라도 같은 숫자가 같은 횟수만큼 있는지."""
    def _key(t):
        try:
            return -abs(float(t.rstrip("%")))
        except ValueError:
            return 0
    missing, counts = [], []
    for t in sorted(set(a) | set(b), key=_key):
        if a[t] == b[t]:
            continue
        item = {"where": f"{where} 숫자 {t}", "made": f"{a[t]}번", "ref": f"{b[t]}번"}
        if a[t] == 0 or b[t] == 0:  # 한쪽에 아예 없는 숫자 - 금액이 빠졌거나 틀렸을 수 있다
            item["kind"] = "숫자 한쪽에만"
            missing.append(item)
        else:  # 같은 숫자가 나오는 횟수만 다름 - 문장 차이일 때가 많다
            item["kind"] = "숫자 횟수"
            counts.append(item)
    out += missing + counts


def short(s, n=70):
    s = (s or "").replace("\n", "⏎")
    return s if len(s) <= n else s[: n - 1] + "…"


# ── 워드 ────────────────────────────────────────────────────
def docx_parts(path: Path):
    import docx  # python-docx
    d = docx.Document(str(path))
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    tables = [[[c.text.strip() for c in row.cells] for row in t.rows] for t in d.tables]
    heads = []
    for s in d.sections:
        for part in (s.header, s.footer):
            try:
                heads += [p.text.strip() for p in part.paragraphs if p.text.strip()]
            except Exception:
                pass
    return {"문단": paras, "표": tables, "머리글·바닥글": heads}


# ── PPT ─────────────────────────────────────────────────────
def pptx_parts(path: Path):
    from pptx import Presentation
    prs = Presentation(str(path))
    slides = []
    for i, sl in enumerate(prs.slides, 1):
        shapes = []
        for sh in sl.shapes:
            item = {"name": sh.name, "text": "", "table": None, "chart": None}
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                item["text"] = "\n".join(p.text for p in sh.text_frame.paragraphs).strip()
            if getattr(sh, "has_table", False) and sh.has_table:
                item["table"] = [[c.text.strip() for c in row.cells] for row in sh.table.rows]
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                data = {}
                try:
                    cats = [str(c) for c in ch.plots[0].categories]
                except Exception:
                    cats = []
                for plot in ch.plots:
                    for s in plot.series:
                        data[s.name] = [None if v is None else round(float(v), 6) for v in s.values]
                item["chart"] = {"categories": cats, "series": data}
            if item["text"] or item["table"] or item["chart"]:
                shapes.append(item)
        notes = ""
        if sl.has_notes_slide:
            notes = sl.notes_slide.notes_text_frame.text.strip()
        slides.append({"no": i, "shapes": shapes, "notes": notes})
    return slides


# ── 비교 ────────────────────────────────────────────────────
def diff_lists(a: list[str], b: list[str], where: str, out: list):
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        if tag == "replace":
            for k in range(max(i2 - i1, j2 - j1)):
                x = a[i1 + k] if i1 + k < i2 else None
                y = b[j1 + k] if j1 + k < j2 else None
                out.append({"where": where, "kind": "바뀜" if x is not None and y is not None else ("없음" if y is None else "추가"),
                            "made": x, "ref": y})
        elif tag == "delete":
            for k in range(i1, i2):
                out.append({"where": where, "kind": "만든 쪽만", "made": a[k], "ref": None})
        elif tag == "insert":
            for k in range(j1, j2):
                out.append({"where": where, "kind": "완성본만", "made": None, "ref": b[k]})


def diff_table(ta, tb, where: str, out: list):
    if ta is None or tb is None:
        out.append({"where": where, "kind": "표 없음", "made": "있음" if ta else "없음", "ref": "있음" if tb else "없음"})
        return
    if (len(ta), max((len(r) for r in ta), default=0)) != (len(tb), max((len(r) for r in tb), default=0)):
        out.append({"where": where, "kind": "표 크기", "made": f"{len(ta)}행", "ref": f"{len(tb)}행"})
    for r in range(min(len(ta), len(tb))):
        for c in range(min(len(ta[r]), len(tb[r]))):
            if ta[r][c] != tb[r][c]:
                out.append({"where": f"{where} {r + 1}행 {c + 1}열", "kind": "표 칸", "made": ta[r][c], "ref": tb[r][c]})


def compare_docx(made: Path, ref: Path, numbers: bool) -> list:
    a, b = docx_parts(made), docx_parts(ref)
    out = []
    if numbers:
        na = numbers_in("\n".join(a["문단"] + [c for t in a["표"] for r in t for c in r]))
        nb = numbers_in("\n".join(b["문단"] + [c for t in b["표"] for r in t for c in r]))
        diff_bags(na, nb, "문서 전체", out)
        return out
    diff_lists(a["문단"], b["문단"], "문단", out)
    diff_lists(a["머리글·바닥글"], b["머리글·바닥글"], "머리글·바닥글", out)
    if len(a["표"]) != len(b["표"]):
        out.append({"where": "표", "kind": "표 수", "made": len(a["표"]), "ref": len(b["표"])})
    for i in range(min(len(a["표"]), len(b["표"]))):
        diff_table(a["표"][i], b["표"][i], f"표{i + 1}", out)
    return out


def compare_pptx(made: Path, ref: Path, numbers: bool) -> list:
    a, b = pptx_parts(made), pptx_parts(ref)
    out = []
    if len(a) != len(b):
        out.append({"where": "슬라이드", "kind": "쪽 수", "made": len(a), "ref": len(b)})
    for sa, sb in zip(a, b):
        w = f"{sa['no']}쪽"
        if numbers:
            ta = "\n".join(s["text"] for s in sa["shapes"]) + "\n" + "\n".join(
                c for s in sa["shapes"] if s["table"] for r in s["table"] for c in r)
            tb = "\n".join(s["text"] for s in sb["shapes"]) + "\n" + "\n".join(
                c for s in sb["shapes"] if s["table"] for r in s["table"] for c in r)
            diff_bags(numbers_in(ta), numbers_in(tb), w, out)
        else:
            diff_lists([s["text"] for s in sa["shapes"] if s["text"]], [s["text"] for s in sb["shapes"] if s["text"]],
                       f"{w} 글자", out)
            tables_a = [s["table"] for s in sa["shapes"] if s["table"]]
            tables_b = [s["table"] for s in sb["shapes"] if s["table"]]
            if len(tables_a) != len(tables_b):
                out.append({"where": f"{w} 표", "kind": "표 수", "made": len(tables_a), "ref": len(tables_b)})
            for i in range(min(len(tables_a), len(tables_b))):
                diff_table(tables_a[i], tables_b[i], f"{w} 표{i + 1}", out)
        charts_a = [s["chart"] for s in sa["shapes"] if s["chart"]]
        charts_b = [s["chart"] for s in sb["shapes"] if s["chart"]]
        if len(charts_a) != len(charts_b):
            out.append({"where": f"{w} 차트", "kind": "차트 수", "made": len(charts_a), "ref": len(charts_b)})
        for i in range(min(len(charts_a), len(charts_b))):
            ca, cb = charts_a[i], charts_b[i]
            if ca["categories"] != cb["categories"]:
                out.append({"where": f"{w} 차트{i + 1} 항목", "kind": "차트", "made": " · ".join(ca["categories"]),
                            "ref": " · ".join(cb["categories"])})
            for name in sorted(set(ca["series"]) | set(cb["series"])):
                va, vb = ca["series"].get(name), cb["series"].get(name)
                if va != vb:
                    out.append({"where": f"{w} 차트{i + 1} 계열 '{name}'", "kind": "차트 값", "made": va, "ref": vb})
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="워드 · PPT 비교")
    ap.add_argument("made")
    ap.add_argument("ref")
    ap.add_argument("--numbers", action="store_true", help="글 속 숫자 모음만 비교(순서 · 문장 무관, 같은 숫자가 같은 횟수만큼)")
    ap.add_argument("--max", type=int, default=40)
    ap.add_argument("--json", default="")
    a = ap.parse_args(argv)
    pm, pr = Path(a.made), Path(a.ref)
    for p in (pm, pr):
        if not p.exists():
            print(f"[오류] 파일이 없습니다: {p}", file=sys.stderr)
            return 2
    ext = pm.suffix.lower()
    if ext != pr.suffix.lower() or ext not in (".docx", ".pptx"):
        print("[오류] 두 파일 모두 .docx 이거나 모두 .pptx 여야 합니다", file=sys.stderr)
        return 2
    try:
        diffs = compare_docx(pm, pr, a.numbers) if ext == ".docx" else compare_pptx(pm, pr, a.numbers)
    except ImportError as e:
        print(f"[오류] {e} - pip install python-docx python-pptx", file=sys.stderr)
        return 2
    except Exception as e:
        print(f"[오류] {type(e).__name__}: {e}", file=sys.stderr)
        return 2
    mode = "숫자만" if a.numbers else "글자 · 표 · 차트"
    print(f"[문서 맞추기 - {mode}] 만든 파일: {pm}")
    print(f"                          완성본  : {pr}")
    print(f"차이 {len(diffs)}개")
    for d in diffs[: a.max]:
        print(f"  ✗ [{d['kind']}] {d['where']}  만든 {short(str(d['made']))} · 완성본 {short(str(d['ref']))}")
    if len(diffs) > a.max:
        print(f"  … 나머지 {len(diffs) - a.max}개는 --json 으로")
    if a.json:
        Path(a.json).parent.mkdir(parents=True, exist_ok=True)
        Path(a.json).write_text(json.dumps(diffs, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    return 0 if not diffs else 1


if __name__ == "__main__":
    sys.exit(main())
