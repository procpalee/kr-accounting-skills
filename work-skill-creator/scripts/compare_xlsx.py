"""두 엑셀 파일을 칸마다 비교한다 - 맞추기(기본) · 기간 비교(--period).

맞추기 : 코드로 다시 만든 파일 vs 사람이 만든 완성본. 목표는 차이 0.
  python -X utf8 compare_xlsx.py <만든.xlsx> <완성본.xlsx> [--formats] [--max 40] [--json 차이.json]

기간 비교 : 그 전 기간 완성본 vs 지난 기간 완성본. 사람이 기간마다 바꾸는 칸(입력 후보)과
           그대로 두는 칸(고정)을 나눠 보여 준다 - 업무를 파악할 때 쓴다.
  python -X utf8 compare_xlsx.py --period <그전.xlsx> <지난.xlsx>

맞추기에서 보는 것
  시트 목록·순서 / 칸마다 값(계산값) · 수식 글자 · 값↔수식(한쪽만 수식 = 값을 붙여 넣은 흔적)
  --formats 를 주면 숫자 서식 · 글꼴 · 채우기 · 테두리 · 맞춤 · 열 너비 · 병합 · 고정한 행 높이 · 숨긴 행 열 · 탭 색 · 틀 고정까지

옵션
  --sheets A,B          이 시트만 비교
  --ignore 범위목록      무시할 칸(쉼표로): 시트!A1:C3 · 시트!F:F · 시트!3:3 · 시트(통째로 - 한쪽에만 있는 시트도). 이유는 스킬 문서에 적는다
  --ignore-file 파일     무시할 범위 목록 파일(업무 스킬 references/known_diffs.txt - 한 줄에 '범위 | 이유 | 영향', # 은 설명)
  --tol 수              숫자 허용 오차(기본 0 - 소수점 끝자리 오차만 허용). 1원 미만을 무시하려면 --tol 0.5
  --blank-zero          빈칸과 0 을 같은 것으로 본다
  --formula-ok          완성본은 값인데 만든 쪽은 수식이고 계산값이 같으면 차이로 세지 않는다(원칙대로 수식으로 바꾼 칸 - 개수만 알림)
  --loose-text          글자 앞뒤 공백·연속 공백 차이 무시
  --loose-formulas      수식의 공백·$ 표시·대소문자 차이 무시
  --recalc              계산값이 비어 있으면 임시 사본을 엑셀로 다시 계산해서 읽는다(윈도우 + 엑셀)
  --max N               화면에 보여 줄 차이 수(기본 40, 전체는 --json)
  --json 파일            전체 결과를 JSON 으로 저장
메모 · 조건부 서식 · 데이터 유효성 · 창 위치 · 시트 선택은 비교하지 않는다(필요하면 눈으로).
종료 코드: 맞추기 = 차이 없음 0 · 차이 있음 1 / 기간 비교 = 늘 0(알려 주기만) / 실행 오류 2
"""
from __future__ import annotations

import argparse
import datetime as dt
import difflib
import json
import re
import shutil
import sys
import tempfile
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

try:
    import openpyxl
    from openpyxl.utils import get_column_letter, range_boundaries
except ImportError:  # pragma: no cover
    print("[오류] openpyxl 이 필요합니다: pip install openpyxl", file=sys.stderr)
    sys.exit(2)

ERRORS = {"#N/A", "#DIV/0!", "#REF!", "#VALUE!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!", "#GETTING_DATA"}
PERIOD_RE = re.compile(r"(20\d{2}[-.]?(0[1-9]|1[0-2]))|(\d{2}\.\s?[1-4]Q)|([1-4]Q)|([1-4]\s?분기)|(상반기|하반기|반기)", re.I)


# ── 읽기 ────────────────────────────────────────────────────
def load_pair(path: Path):
    """(수식 통합문서, 계산값 통합문서)."""
    wf = openpyxl.load_workbook(path, data_only=False)
    wv = openpyxl.load_workbook(path, data_only=True)
    return wf, wv


def cell_map(ws) -> dict:
    """(행, 열) → Cell. 만들어진 칸만(빈 칸 전체를 돌지 않는다)."""
    cells = getattr(ws, "_cells", None)
    if cells is not None:
        return dict(cells)
    return {(c.row, c.column): c for row in ws.iter_rows() for c in row}


def formula_text(v):
    """수식이면 글자, 아니면 None. 배열수식·데이터표도 글자로."""
    if isinstance(v, str) and v.startswith("="):
        return v
    t = getattr(v, "text", None)  # ArrayFormula
    if t is not None:
        ref = getattr(v, "ref", "")
        return f"{{{t}}}[{ref}]"
    if type(v).__name__ == "DataTableFormula":
        return "=TABLE(" + ",".join(str(getattr(v, k, "")) for k in ("r1", "r2")) + ")"
    return None


def norm_value(v, loose_text=False):
    if v is None:
        return None
    if isinstance(v, str):
        if loose_text:
            v = " ".join(v.split())
        return v if v != "" else None
    if isinstance(v, dt.time):
        return v.isoformat()
    return v


def norm_formula(f: str, loose: bool) -> str:
    if not loose:
        return f
    return re.sub(r"\s+", "", f).replace("$", "").upper()


def to_number(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, (dt.datetime, dt.date)):
        from openpyxl.utils.datetime import to_excel
        try:
            return float(to_excel(v))
        except Exception:
            return None
    return None


def same_value(a, b, tol: float) -> bool:
    """숫자(날짜는 엑셀 일련번호로 바꿔서)는 허용 오차 안이면 같다. 부동소수 잡음(1e-9 상대)은 늘 허용."""
    if a is None and b is None:
        return True
    na, nb = to_number(a), to_number(b)
    if na is not None and nb is not None:
        lim = max(tol, 1e-9 * max(abs(na), abs(nb)), 1e-9)
        return abs(na - nb) <= lim
    return a == b


def show(v, limit=60) -> str:
    if v is None:
        return "(빈칸)"
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, int):
        return f"{v:,}"
    if isinstance(v, float):
        return f"{v:,.6f}".rstrip("0").rstrip(".")
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).replace("\n", "⏎")
    return s if len(s) <= limit else s[: limit - 1] + "…"


# ── 무시 범위 ────────────────────────────────────────────────
def read_ignore_file(path: str) -> list[str]:
    """known_diffs.txt - 한 줄에 '범위 | 이유 | 영향'. # 로 시작하는 줄 · 빈 줄은 건너뛴다. 범위만 돌려준다."""
    if not path:
        return []
    out = []
    for line in Path(path).read_text(encoding="utf-8-sig").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        rng = line.split("|", 1)[0].strip()
        if rng:
            out.append(rng)
    return out


IGNORE_MODES = {"값만:": "formula", "서식만:": "format"}  # 앞에 붙이면: 값만 = 수식 차이는 허용하고 값은 비교 / 서식만 = 서식만 무시


def parse_ignore(spec, extra: list[str] | None = None) -> dict:
    """{'all' | 'formula' | 'format': {시트: [(min_col, min_row, max_col, max_row) | None(통째로)]}}."""
    rules = {"all": defaultdict(list), "formula": defaultdict(list), "format": defaultdict(list)}
    parts = [p.strip() for p in (spec or "").split(",") if p.strip()] + list(extra or [])
    for part in parts:
        mode = "all"
        for pre, m in IGNORE_MODES.items():
            if part.startswith(pre):
                mode, part = m, part[len(pre):].strip()
                break
        if "!" in part:
            sheet, rng = part.rsplit("!", 1)
            sheet = sheet.strip().strip("'")
            try:
                b = range_boundaries(rng.strip().replace("$", ""))
            except ValueError:
                raise SystemExit(f"[오류] 무시 범위를 읽지 못했습니다: {part}")
            rules[mode][sheet].append(b)
        else:
            rules[mode][part.strip().strip("'")].append(None)
    return rules


def sheet_ignored(rules, sheet) -> bool:
    return any(b is None for b in rules.get(sheet, []))


def ignored(rules, sheet, row, col) -> bool:
    for b in rules.get(sheet, []):
        if b is None:
            return True
        c1, r1, c2, r2 = b
        if (c1 is None or c1 <= col) and (c2 is None or col <= c2) and (r1 is None or r1 <= row) and (r2 is None or row <= r2):
            return True
    return False


# ── 서식 ────────────────────────────────────────────────────
def color_key(c):
    if c is None:
        return None
    t = getattr(c, "type", None)
    tint = getattr(c, "tint", 0) or 0
    if t == "rgb":
        v = c.rgb if isinstance(c.rgb, str) else None
        if v in (None, "00000000"):
            return None
        return v.upper()
    if t == "theme":
        return f"theme{c.theme}" + (f"{tint:+.2f}" if tint else "")
    if t == "indexed":
        return None if c.indexed in (64,) else f"idx{c.indexed}"
    return None


def style_sig(cell) -> dict:
    f, fl, b, al = cell.font, cell.fill, cell.border, cell.alignment
    fill = None
    if fl is not None and getattr(fl, "fill_type", None) not in (None, "none"):
        fill = f"{fl.fill_type}:{color_key(fl.fgColor)}"
    return {
        "숫자서식": cell.number_format,
        "굵게": bool(f.b), "기울임": bool(f.i), "밑줄": f.u or None,
        "글자색": color_key(f.color), "글자크기": f.sz, "글꼴": f.name,
        "채우기": fill,
        "테두리": "/".join(str(getattr(getattr(b, side), "style", None) or "-") for side in ("left", "right", "top", "bottom")),
        "가로맞춤": al.horizontal, "줄바꿈": bool(al.wrap_text),
    }


def col_widths(ws) -> dict:
    out = {}
    for key, dim in ws.column_dimensions.items():
        lo = getattr(dim, "min", None) or openpyxl.utils.column_index_from_string(key)
        hi = getattr(dim, "max", None) or lo
        for c in range(lo, hi + 1):
            if dim.width is not None:
                out[c] = round(float(dim.width), 2)
    return out


def tab_color(ws):
    return color_key(getattr(ws.sheet_properties, "tabColor", None))


def raw_rows(path: Path) -> dict:
    """{시트 이름: {행: (높이 | None, 고정 여부, 숨김 여부)}} - 파일 원문(XML)의 <row ht customHeight hidden> 그대로."""
    import zipfile
    import xml.etree.ElementTree as ET
    m = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rel = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    out = {}
    try:
        with zipfile.ZipFile(path) as z:
            wb = ET.fromstring(z.read("xl/workbook.xml"))
            rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
            target = {r.get("Id"): r.get("Target") for r in rels}
            sheets = wb.find(f"{m}sheets")
            for sh in (sheets if sheets is not None else []):
                t = target.get(sh.get(rel), "")
                inner = t.lstrip("/") if t.startswith("/") else "xl/" + t
                rows = {}
                try:
                    with z.open(inner) as f:
                        for _, el in ET.iterparse(f, events=("end",)):
                            if el.tag == f"{m}row":
                                ht, cust, hid = el.get("ht"), el.get("customHeight"), el.get("hidden")
                                custom = cust in ("1", "true")
                                hidden = hid in ("1", "true")
                                if ht or custom or hidden:
                                    rows[int(el.get("r"))] = (float(ht) if ht else None, custom, hidden)
                                el.clear()
                except KeyError:
                    pass
                out[sh.get("name")] = rows
    except Exception:
        return {}
    return out


# ── 맞추기 ──────────────────────────────────────────────────
BLANK_STYLE = {"채우기": None, "테두리": "-/-/-/-"}  # 한쪽에 칸이 없을 때 그쪽을 이 서식으로 보고 비교(눈에 띄는 것만)


def compare(made: Path, ref: Path, a) -> dict:
    mf, mv = load_pair(made)
    rf, rv = load_pair(ref)
    rules = parse_ignore(a.ignore, read_ignore_file(getattr(a, "ignore_file", "")))
    want = [s.strip() for s in a.sheets.split(",")] if a.sheets else None
    blank_zero = getattr(a, "blank_zero", False)
    formula_ok = getattr(a, "formula_ok", False)
    r_all, r_formula, r_format = rules["all"], rules["formula"], rules["format"]

    diffs = []
    notes = []
    formula_ok_count = 0
    value_only_count = 0  # known_diffs '값만:' - 수식 차이는 허용하고 값만 비교한 칸
    m_names, r_names = mf.sheetnames, rf.sheetnames
    only_m = [s for s in m_names if s not in r_names and not sheet_ignored(r_all, s)]
    only_r = [s for s in r_names if s not in m_names and not sheet_ignored(r_all, s)]
    common = [s for s in r_names if s in m_names and not sheet_ignored(r_all, s)]
    if want:
        missing = [s for s in want if s not in common]
        if missing:
            notes.append(f"양쪽에 다 있지 않은(또는 무시한) 시트: {', '.join(missing)}")
        common = [s for s in common if s in want]
    else:
        for s in only_m:
            diffs.append({"sheet": s, "cell": "", "kind": "시트", "made": "있음", "ref": "없음", "note": "만든 파일에만 있는 시트"})
        for s in only_r:
            diffs.append({"sheet": s, "cell": "", "kind": "시트", "made": "없음", "ref": "있음", "note": "완성본에만 있는 시트"})
        order_m = [s for s in m_names if s in r_names and not sheet_ignored(r_all, s)]
        order_r = [s for s in r_names if s in m_names and not sheet_ignored(r_all, s)]
        if order_m != order_r:
            diffs.append({"sheet": "", "cell": "", "kind": "시트", "made": " · ".join(order_m), "ref": " · ".join(order_r),
                          "note": "시트 순서가 다름"})

    compared = 0
    missing_calc = {"made": [0, 0], "ref": [0, 0]}  # [수식 칸, 계산값 없는 수식 칸]
    err_cells = []
    raw_m, raw_r = (raw_rows(made), raw_rows(ref)) if a.formats else ({}, {})
    for s in common:
        wmf, wmv, wrf, wrv = mf[s], mv[s], rf[s], rv[s]
        cm_f, cm_v, cr_f, cr_v = cell_map(wmf), cell_map(wmv), cell_map(wrf), cell_map(wrv)
        keys = set(cm_f) | set(cr_f)
        for (r, c) in sorted(keys):
            if ignored(r_all, s, r, c):
                continue
            only_values = ignored(r_formula, s, r, c)
            no_format = ignored(r_format, s, r, c)
            cmf, crf = cm_f.get((r, c)), cr_f.get((r, c))
            fm = formula_text(cmf.value) if cmf is not None else None
            fr = formula_text(crf.value) if crf is not None else None
            vm = norm_value(cm_v[(r, c)].value if (r, c) in cm_v else None, a.loose_text)
            vr = norm_value(cr_v[(r, c)].value if (r, c) in cr_v else None, a.loose_text)
            if fm is None and cmf is not None:
                vm = norm_value(cmf.value, a.loose_text)
            if fr is None and crf is not None:
                vr = norm_value(crf.value, a.loose_text)
            if blank_zero:
                vm = 0 if vm is None and fm is None else vm
                vr = 0 if vr is None and fr is None else vr
            has_content = any(x is not None and not (blank_zero and x == 0) for x in (fm, fr, vm, vr))
            style_diff = None
            if a.formats and not no_format and (cmf is not None or crf is not None):
                sm = style_sig(cmf) if cmf is not None else None
                sr = style_sig(crf) if crf is not None else None
                if sm is not None and sr is not None:
                    bad = [k for k in sm if sm[k] != sr[k]]
                else:  # 한쪽에만 칸이 있으면 다른 쪽은 빈 서식으로 보고 눈에 띄는 것(채우기 · 테두리)만
                    one = sm if sm is not None else sr
                    sm = sm if sm is not None else {**one, **BLANK_STYLE}
                    sr = sr if sr is not None else {**one, **BLANK_STYLE}
                    bad = [k for k in BLANK_STYLE if sm[k] != sr[k]]
                if bad:
                    style_diff = "; ".join(f"{k} {show(sm[k], 24)}≠{show(sr[k], 24)}" for k in bad)
            if not has_content and style_diff is None:
                continue
            compared += 1
            addr = f"{get_column_letter(c)}{r}"
            if fm is not None:
                missing_calc["made"][0] += 1
                if vm is None:
                    missing_calc["made"][1] += 1
            if fr is not None:
                missing_calc["ref"][0] += 1
                if vr is None:
                    missing_calc["ref"][1] += 1
            if isinstance(vm, str) and vm in ERRORS:
                err_cells.append(f"{s}!{addr} {vm}")
            vals_same = same_value(vm, vr, a.tol)
            formula_differs = ((fm is None) != (fr is None)) or (
                fm is not None and norm_formula(fm, a.loose_formulas) != norm_formula(fr, a.loose_formulas))
            if only_values and formula_differs:
                if vals_same:
                    value_only_count += 1  # 수식 차이는 known_diffs 로 허용 - 값은 같다
                else:
                    diffs.append({"sheet": s, "cell": addr, "kind": "값", "made": vm, "ref": vr,
                                  "note": "수식 차이는 허용한 칸(값만:)인데 값이 다름"})
            elif formula_ok and fm is not None and fr is None and vals_same:
                formula_ok_count += 1  # 사람이 친 값을 원칙대로 수식으로 바꾼 칸 - 값이 같으면 차이로 세지 않는다
            elif (fm is None) != (fr is None):
                diffs.append({"sheet": s, "cell": addr, "kind": "값↔수식",
                              "made": fm if fm is not None else vm, "ref": fr if fr is not None else vr,
                              "note": ("값은 같음 - " if vals_same else "") +
                                      ("만든 쪽이 값(붙여 넣음?)" if fm is None else "완성본 쪽이 값(사람이 직접 입력)")})
            elif fm is not None and norm_formula(fm, a.loose_formulas) != norm_formula(fr, a.loose_formulas):
                diffs.append({"sheet": s, "cell": addr, "kind": "수식", "made": fm, "ref": fr,
                              "note": "값은 같음" if vals_same else f"값도 다름: {show(vm)} ≠ {show(vr)}"})
            elif not vals_same:
                n = ""
                nm, nr = to_number(vm), to_number(vr)
                if nm is not None and nr is not None and not isinstance(vm, bool):
                    n = f"차이 {show(nm - nr)}"
                diffs.append({"sheet": s, "cell": addr, "kind": "값", "made": vm, "ref": vr, "note": n})
            if style_diff:
                diffs.append({"sheet": s, "cell": addr, "kind": "서식", "made": "", "ref": "", "note": style_diff})
        if a.formats:
            wm, wr = col_widths(wmf), col_widths(wrf)
            used_cols = {c for (_, c) in keys}
            for c in sorted(used_cols):
                if all(ignored(rr, s, row, c) for row in (1, 10 ** 6) for rr in (r_all,)) or \
                        all(ignored(r_format, s, row, c) for row in (1, 10 ** 6)):
                    continue
                if wm.get(c) != wr.get(c):
                    diffs.append({"sheet": s, "cell": get_column_letter(c), "kind": "서식", "made": wm.get(c), "ref": wr.get(c),
                                  "note": "열 너비"})
            if sheet_ignored(r_format, s):
                continue
            mm = {str(x) for x in wmf.merged_cells.ranges}
            mr = {str(x) for x in wrf.merged_cells.ranges}
            for x in sorted(mm - mr):
                diffs.append({"sheet": s, "cell": x, "kind": "서식", "made": "병합", "ref": "", "note": "만든 쪽만 병합"})
            for x in sorted(mr - mm):
                diffs.append({"sheet": s, "cell": x, "kind": "서식", "made": "", "ref": "병합", "note": "완성본 쪽만 병합"})
            # 고정 행 높이 · 숨긴 행 - 파일 원문(XML)의 customHeight 로 본다. openpyxl 은 높이 값만 있어도 '고정'으로 보여 주므로 쓰지 않는다.
            # (엑셀이 기록한 자동 높이는 비교하지 않는다. openpyxl 로 다시 저장하면 그 자동 높이가 고정으로 굳는 일을 잡는 게 목적)
            rows_m, rows_r = raw_m.get(s, {}), raw_r.get(s, {})
            for r in sorted(set(rows_m) | set(rows_r)):
                if (ignored(r_all, s, r, 1) and ignored(r_all, s, r, 16384)) or \
                        (ignored(r_format, s, r, 1) and ignored(r_format, s, r, 16384)):
                    continue
                hm_, cm_, xm_ = rows_m.get(r, (None, False, False))
                hr_, cr_, xr_ = rows_r.get(r, (None, False, False))
                fm_ = round(hm_, 2) if cm_ and hm_ else None
                fr_ = round(hr_, 2) if cr_ and hr_ else None
                if fm_ != fr_:
                    diffs.append({"sheet": s, "cell": f"{r}행", "kind": "서식", "made": fm_ or "자동", "ref": fr_ or "자동",
                                  "note": "행 높이(고정)"})
                if xm_ != xr_:
                    diffs.append({"sheet": s, "cell": f"{r}행", "kind": "서식", "made": "숨김" if xm_ else "",
                                  "ref": "숨김" if xr_ else "", "note": "숨긴 행"})
            hid_m = {k for k, d in wmf.column_dimensions.items() if d.hidden}
            hid_r = {k for k, d in wrf.column_dimensions.items() if d.hidden}
            for k in sorted(hid_m ^ hid_r):
                diffs.append({"sheet": s, "cell": k, "kind": "서식", "made": "숨김" if k in hid_m else "", "ref": "숨김" if k in hid_r else "",
                              "note": "숨긴 열"})
            if tab_color(wmf) != tab_color(wrf):
                diffs.append({"sheet": s, "cell": "", "kind": "서식", "made": tab_color(wmf), "ref": tab_color(wrf), "note": "탭 색"})
            if (wmf.freeze_panes or None) != (wrf.freeze_panes or None):
                diffs.append({"sheet": s, "cell": "", "kind": "서식", "made": wmf.freeze_panes, "ref": wrf.freeze_panes, "note": "틀 고정"})

    for side, (nf, nmiss) in missing_calc.items():
        if nf and nmiss / nf > 0.5:
            who = "만든 파일" if side == "made" else "완성본"
            notes.append(f"{who}의 수식 칸 {nf:,}개 중 {nmiss:,}개에 계산값이 없습니다 - 값 비교가 의미 없습니다. "
                         f"엑셀로 다시 계산해 저장하거나 --recalc 를 주세요.")
    if err_cells:
        notes.append(f"만든 파일에 오류 값 {len(err_cells)}칸: " + ", ".join(err_cells[:8]) + (" …" if len(err_cells) > 8 else ""))
    if value_only_count:
        notes.append(f"수식이 달라도 되는 칸(known_diffs '값만:') {value_only_count:,}개는 값만 비교했고 모두 같았습니다.")
    if formula_ok_count:
        notes.append(f"완성본은 값인데 만든 쪽은 수식이고 계산값이 같은 칸 {formula_ok_count:,}개는 차이로 세지 않았습니다(--formula-ok).")
    ignored_sheets = [s for s in set(m_names) | set(r_names) if sheet_ignored(r_all, s)]
    if ignored_sheets:
        notes.append(f"무시한 시트: {', '.join(sorted(ignored_sheets))}")
    return {"mode": "맞추기", "made": str(made), "ref": str(ref), "sheets": common, "compared": compared,
            "diffs": diffs, "notes": notes, "missing_calc": missing_calc, "formats": bool(a.formats)}


def recalc_copy(path: Path) -> Path:
    from wsc_xl import excel_recalc
    tmp = Path(tempfile.mkdtemp(prefix="wsc_cmp_")) / Path(path).name
    shutil.copy2(path, tmp)
    excel_recalc(tmp)
    return tmp


def print_result(res: dict, a) -> None:
    diffs = res["diffs"]
    print(f"[맞추기] 만든 파일: {res['made']}")
    print(f"         완성본  : {res['ref']}")
    what = "값 · 수식 · 서식" if res.get("formats") else "값 · 수식(서식은 --formats)"
    print(f"비교한 시트 {len(res['sheets'])}개 · 칸 {res['compared']:,}개 ({what}) - 차이 {len(diffs):,}개")
    if diffs:
        kinds = Counter(d["kind"] for d in diffs)
        by_sheet = Counter(d["sheet"] or "(통합문서)" for d in diffs)
        print("  종류별: " + " · ".join(f"{k} {n:,}" for k, n in kinds.most_common()))
        print("  시트별: " + " · ".join(f"{k} {n:,}" for k, n in by_sheet.most_common(12)))
        print()
        for d in diffs[: a.max]:
            where = f"{d['sheet']}!{d['cell']}" if d["cell"] else (d["sheet"] or "(통합문서)")
            if d["kind"] == "서식" and d["made"] == "" and d["ref"] == "":
                print(f"  ✗ [서식] {where}  {d['note']}")
            else:
                note = f"  ({d['note']})" if d["note"] else ""
                print(f"  ✗ [{d['kind']}] {where}  만든 {show(d['made'])} · 완성본 {show(d['ref'])}{note}")
        if len(diffs) > a.max:
            print(f"  … 나머지 {len(diffs) - a.max:,}개는 --json 으로 저장해서 보세요")
    for n in res["notes"]:
        print(f"[주의] {n}")


# ── 기간 비교 ────────────────────────────────────────────────
DIGITS = re.compile(r"\d+")
ALIGN_LIMIT = 6000  # 이보다 행이 많은 시트는 행 맞춤 없이 같은 행끼리 비교(원장처럼 큰 데이터 시트)


def ranges_by_col(coords):
    """[(행, 열)] → 열별로 이어진 행을 묶은 범위 목록 [(시작주소, 끝주소, 개수, 첫 칸)]."""
    by_col = defaultdict(list)
    for r, c in coords:
        by_col[c].append(r)
    out = []
    for c in sorted(by_col):
        rows = sorted(by_col[c])
        start = prev = rows[0]
        for r in rows[1:] + [None]:
            if r is not None and r == prev + 1:
                prev = r
                continue
            out.append((f"{get_column_letter(c)}{start}", f"{get_column_letter(c)}{prev}", prev - start + 1, (start, c)))
            if r is not None:
                start = prev = r
    return out


def row_ranges(rows):
    """[3, 4, 5, 9] → [(3, 5), (9, 9)]."""
    out = []
    rows = sorted(rows)
    if not rows:
        return out
    start = prev = rows[0]
    for r in rows[1:] + [None]:
        if r is not None and r == prev + 1:
            prev = r
            continue
        out.append((start, prev))
        if r is not None:
            start = prev = r
    return out


def sheet_rows(wf, wv) -> dict:
    """{행: {열: (수식 글자 | None, 값)}} - 내용이 있는 칸만."""
    cf, cv = cell_map(wf), cell_map(wv)
    rows = defaultdict(dict)
    for (r, c), cell in cf.items():
        f = formula_text(cell.value)
        v = norm_value(cell.value) if f is None else norm_value(cv[(r, c)].value if (r, c) in cv else None)
        if f is None and v is None:
            continue
        rows[r][c] = (f, v)
    return rows


def row_sig(cells: dict) -> tuple:
    """행의 모양: 글자 칸은 글자 그대로, 숫자·날짜는 종류만, 수식은 숫자(행 번호·기간)를 지운 모양."""
    sig = []
    for c in sorted(cells):
        f, v = cells[c]
        if f is not None:
            sig.append((c, "f", DIGITS.sub("#", f)))
        elif isinstance(v, str):
            sig.append((c, "t", v.strip()))
        elif isinstance(v, bool):
            sig.append((c, "b"))
        elif isinstance(v, (dt.date, dt.datetime)):
            sig.append((c, "d"))
        elif isinstance(v, (int, float)):
            sig.append((c, "n"))
        else:
            sig.append((c, "x"))
    return tuple(sig)


def align_rows(rows_o: dict, rows_n: dict):
    """그 전 기간 행과 지난 기간 행을 짝짓는다. (짝 [(옛 행, 새 행)], 새 행 목록, 없어진 행 목록, 맞춤 여부)."""
    max_o, max_n = max(rows_o, default=0), max(rows_n, default=0)
    if max(max_o, max_n) > ALIGN_LIMIT:
        top = max(max_o, max_n)
        return [(r, r) for r in range(1, top + 1)], [], [], False
    seq_o = [row_sig(rows_o.get(r, {})) for r in range(1, max_o + 1)]
    seq_n = [row_sig(rows_n.get(r, {})) for r in range(1, max_n + 1)]
    sm = difflib.SequenceMatcher(None, seq_o, seq_n, autojunk=False)
    pairs, ins, dels = [], [], []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            pairs += [(i1 + k + 1, j1 + k + 1) for k in range(i2 - i1)]
        elif tag == "replace":  # 모양이 다른 행 - 앞에서부터 같은 수만큼 짝짓고 남는 것은 새 행 / 없어진 행
            n = min(i2 - i1, j2 - j1)
            pairs += [(i1 + k + 1, j1 + k + 1) for k in range(n)]
            dels += [i1 + k + 1 for k in range(n, i2 - i1)]
            ins += [j1 + k + 1 for k in range(n, j2 - j1)]
        elif tag == "delete":
            dels += list(range(i1 + 1, i2 + 1))
        elif tag == "insert":
            ins += list(range(j1 + 1, j2 + 1))
    ins = [r for r in ins if rows_n.get(r)]
    dels = [r for r in dels if rows_o.get(r)]
    return pairs, ins, dels, True


def first_label(cells: dict) -> str:
    for c in sorted(cells):
        f, v = cells[c]
        if f is None and isinstance(v, str) and v.strip():
            return v.strip()
    for c in sorted(cells):
        f, v = cells[c]
        return show(f if f is not None else v, 30)
    return ""


CATS = ("상수 바뀜", "상수→수식", "수식→상수", "수식 구조 바뀜", "수식 숫자만 바뀜", "칸 새로 채움", "칸 비움")


def period_compare(old: Path, new: Path, a) -> dict:
    of, ov = load_pair(old)
    nf, nv = load_pair(new)
    res = {"mode": "기간 비교", "old": str(old), "new": str(new), "sheets": [], "groups": {}, "rows": {}, "notes": []}
    res["sheets_only_old"] = [s for s in of.sheetnames if s not in nf.sheetnames]
    res["sheets_only_new"] = [s for s in nf.sheetnames if s not in of.sheetnames]
    want = [x.strip() for x in a.sheets.split(",")] if a.sheets else None
    cats = {k: defaultdict(list) for k in CATS}
    stats = defaultdict(Counter)
    samples = {}
    for s in [x for x in nf.sheetnames if x in of.sheetnames]:
        if want and s not in want:
            continue
        res["sheets"].append(s)
        ro_map, rn_map = sheet_rows(of[s], ov[s]), sheet_rows(nf[s], nv[s])
        pairs, ins, dels, aligned = align_rows(ro_map, rn_map)
        if not aligned:
            res["notes"].append(f"{s}: 행이 {ALIGN_LIMIT:,}개를 넘어 행 맞춤 없이 같은 행끼리 비교했습니다.")
        stats[s]["밀린 행"] = sum(1 for ro, rn in pairs if ro != rn and (ro_map.get(ro) or rn_map.get(rn)))
        res["rows"][s] = {
            "new": [{"rows": f"{x}" if x == y else f"{x}:{y}", "count": y - x + 1, "label": first_label(rn_map.get(x, {}))}
                    for x, y in row_ranges(ins)],
            "gone": [{"rows": f"{x}" if x == y else f"{x}:{y}", "count": y - x + 1, "label": first_label(ro_map.get(x, {}))}
                     for x, y in row_ranges(dels)],
        }
        for ro, rn in pairs:
            co, cn = ro_map.get(ro, {}), rn_map.get(rn, {})
            for c in set(co) | set(cn):
                fo, vo = co.get(c, (None, None))
                fn, vn = cn.get(c, (None, None))
                o_has, n_has = (fo is not None or vo is not None), (fn is not None or vn is not None)
                if o_has and not n_has:
                    cat, key = "칸 비움", (ro, c)
                elif n_has and not o_has:
                    cat, key = "칸 새로 채움", (rn, c)
                elif fo is None and fn is None:
                    if same_value(vo, vn, a.tol):
                        stats[s]["상수 그대로"] += 1
                        continue
                    cat, key = "상수 바뀜", (rn, c)
                elif fo is None:
                    cat, key = "상수→수식", (rn, c)
                elif fn is None:
                    cat, key = "수식→상수", (rn, c)
                elif fo == fn:
                    stats[s]["수식 그대로"] += 1
                    continue
                elif DIGITS.sub("#", fo) == DIGITS.sub("#", fn):
                    cat, key = "수식 숫자만 바뀜", (rn, c)
                else:
                    cat, key = "수식 구조 바뀜", (rn, c)
                stats[s][cat] += 1
                cats[cat][s].append(key)
                samples[(s, cat, key)] = (fo if fo is not None else vo, fn if fn is not None else vn, ro, rn)
    for cat, per_sheet in cats.items():
        groups = []
        for s, coords in per_sheet.items():
            by_hint = defaultdict(list)  # 같은 열에서 이어져도 성격(날짜 · 숫자 · 글 · 수식 변화 종류)이 다르면 따로 묶는다
            for k in coords:
                o, nn, _, _ = samples[(s, cat, k)]
                by_hint[hint(cat, o, nn)].append(k)
            for h, hcoords in by_hint.items():
                for start, end, n, first in ranges_by_col(hcoords):
                    o, nn, ro, rn = samples[(s, cat, first)]
                    moved = f"그 전 {ro}행" if (ro != rn and cat != "칸 비움") else ""
                    groups.append({"sheet": s, "range": start if start == end else f"{start}:{end}", "count": n,
                                   "old": o, "new": nn, "moved": moved, "hint": h})
        groups.sort(key=lambda g: (-g["count"], g["sheet"], g["range"]))
        res["groups"][cat] = groups
    res["stats"] = {s: dict(c) for s, c in stats.items()}
    return res


SHEET_REF = re.compile(r"'([^']+)'!|([A-Za-z0-9_.가-힣]+)!")
CELL_REF = re.compile(r"\$?[A-Z]{1,3}\$?\d+")
PERIOD_NUM = re.compile(r"^(20\d{2}(0[1-9]|1[0-2])([0-3]\d)?|20\d{2})$")


def formula_change(fo: str, fn: str) -> dict:
    """숫자만 바뀐 두 수식에서 무엇이 바뀌었나: 시트 이름 · 칸 주소(행) · 수식 안 상수."""
    def sheets(f):
        return [m.group(1) or m.group(2) for m in SHEET_REF.finditer(f)]

    def cells(f):
        return CELL_REF.findall(SHEET_REF.sub("", f))

    def consts(f):
        return re.findall(r"\d+", CELL_REF.sub("", SHEET_REF.sub("", f)))
    co, cn = consts(fo), consts(fn)
    changed_consts = [x for x, y in zip(co, cn) if x != y] + [y for x, y in zip(co, cn) if x != y]
    return {"sheet": sheets(fo) != sheets(fn), "rows": cells(fo) != cells(fn), "const": co != cn,
            "period_const": any(PERIOD_NUM.match(x) for x in changed_consts)}


def hint(cat, o, n) -> str:
    """바뀐 칸이 무엇인지 짐작."""
    if cat == "수식 숫자만 바뀜":
        ch = formula_change(str(o), str(n))
        if ch["period_const"]:
            return "수식 안에 박힌 기간 숫자가 바뀜 - 기준일(기준 년월) 칸을 두고 그 칸을 참조하게 바꿀 후보"
        if ch["const"]:
            return "수식 안 상수(요율 · 금액)가 바뀜 - 판단값이면 입력 칸으로 빼낼 후보"
        if ch["sheet"] and ch["rows"]:
            return "기간 시트 이름과 행 번호가 바뀜 - 새 기간 시트를 보게 하고 행 수에 맞춰 범위를 다시 씀"
        if ch["sheet"]:
            return "기간이 든 시트 참조만 바뀜 - 코드가 새 기간 시트 이름으로 바꿔 씀"
        return "범위 · 행 번호만 바뀜 - 행 수 변화를 따라감(코드가 행 수에 맞춰 범위를 다시 씀)"
    if cat in ("상수 바뀜", "칸 새로 채움"):
        if isinstance(n, (dt.date, dt.datetime)):
            return "날짜 - 기준일이면 한 칸에서만 바꾸고 나머지는 그 칸을 참조하는 수식으로"
        if isinstance(n, str) and PERIOD_RE.search(n):
            return "기간이 든 글 - 기준일 칸을 참조하는 수식으로 바꿀 후보"
        if isinstance(n, (int, float)) and not isinstance(n, bool):
            return "숫자 - 원자료 · 이월(지난 기간 결과) · 판단값 중 무엇인지 확인"
    if cat == "수식→상수":
        return "수식을 값으로 덮어쓴 흔적일 수 있음 - 사람이 손으로 고친 칸인지 확인"
    return ""


def print_period(res: dict, a) -> None:
    print(f"[기간 비교] 그 전 기간: {res['old']}")
    print(f"            지난 기간: {res['new']}")
    if res["sheets_only_new"]:
        print("새 시트(지난 기간에만): " + ", ".join(res["sheets_only_new"]))
    if res["sheets_only_old"]:
        print("없어진 시트(그 전 기간에만): " + ", ".join(res["sheets_only_old"]))
    print()
    print("시트별 - 그대로 둔 칸(고정 후보) / 바뀐 칸(입력 · 구조 후보) / 행")
    for s in res["sheets"]:
        st = res["stats"].get(s, {})
        same = st.get("상수 그대로", 0) + st.get("수식 그대로", 0)
        changed = sum(v for k, v in st.items() if k in CATS)
        r = res["rows"].get(s, {})
        nnew = sum(x["count"] for x in r.get("new", []))
        ngone = sum(x["count"] for x in r.get("gone", []))
        print(f"  {s}: 그대로 {same:,} (상수 {st.get('상수 그대로', 0):,} · 수식 {st.get('수식 그대로', 0):,})"
              f" / 바뀜 {changed:,} / 새 행 {nnew} · 없어진 행 {ngone} · 밀린 행 {st.get('밀린 행', 0)}")
    row_lines = []
    for s in res["sheets"]:
        for x in res["rows"].get(s, {}).get("new", []):
            row_lines.append(f"  + {s}!{x['rows']}행 ({x['count']}행)  {show(x['label'], 40)}")
        for x in res["rows"].get(s, {}).get("gone", []):
            row_lines.append(f"  - {s}!{x['rows']}행 (그 전 파일, {x['count']}행)  {show(x['label'], 40)}")
    if row_lines:
        print("\n■ 늘어난 행(+) · 없어진 행(-) - 명단 · 명세처럼 기간마다 행 수가 바뀌는 곳")
        for line in row_lines[: a.max]:
            print(line)
        if len(row_lines) > a.max:
            print(f"  … 나머지 {len(row_lines) - a.max:,}줄은 --json 으로")
    titles = {
        "상수 바뀜": "사람이 기간마다 넣은 값 - 원자료 · 이월 · 판단값 · 기간 문구 후보",
        "칸 새로 채움": "같은 행에 새로 채운 칸",
        "상수→수식": "값이었다가 수식이 된 칸",
        "수식→상수": "수식이었다가 값이 된 칸",
        "수식 구조 바뀜": "수식 모양이 바뀐 칸 - 계산 방법이 바뀌었는지 확인",
        "수식 숫자만 바뀜": "수식에서 숫자(행 번호 · 기간)만 바뀐 칸 - 대개 행 수 변화를 따라간 것",
        "칸 비움": "같은 행에서 비운 칸",
    }
    for cat, title in titles.items():
        groups = res["groups"].get(cat) or []
        if not groups:
            continue
        total = sum(g["count"] for g in groups)
        print(f"\n■ {title} ({total:,}칸, {len(groups):,}묶음)")
        for g in groups[: a.max]:
            h = f"  ← {g['hint']}" if g["hint"] else ""
            mv = f" [{g['moved']}]" if g["moved"] else ""
            print(f"  {g['sheet']}!{g['range']} ({g['count']}칸){mv}  예: {show(g['old'], 40)} → {show(g['new'], 40)}{h}")
        if len(groups) > a.max:
            print(f"  … 나머지 {len(groups) - a.max:,}묶음은 --json 으로")
    for n in res["notes"]:
        print(f"[주의] {n}")


# ── 실행 ────────────────────────────────────────────────────
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="두 엑셀 파일 칸마다 비교(맞추기 · 기간 비교)")
    ap.add_argument("first", help="맞추기: 만든 파일 / --period: 그 전 기간 완성본")
    ap.add_argument("second", help="맞추기: 완성본 / --period: 지난 기간 완성본")
    ap.add_argument("--period", action="store_true", help="기간 비교 모드")
    ap.add_argument("--formats", action="store_true", help="서식 · 열 너비 · 병합 · 탭 색 · 틀 고정까지 비교")
    ap.add_argument("--sheets", default="", help="비교할 시트(쉼표)")
    ap.add_argument("--ignore", default="", help="무시할 범위(쉼표): 시트!A1:C3, 시트!F:F, 시트")
    ap.add_argument("--ignore-file", default="", help="무시할 범위 목록 파일(known_diffs.txt)")
    ap.add_argument("--tol", type=float, default=0.0, help="숫자 허용 오차(기본 0)")
    ap.add_argument("--blank-zero", action="store_true", help="빈칸과 0 을 같게 봄")
    ap.add_argument("--formula-ok", action="store_true", help="완성본 값 ↔ 만든 쪽 수식(계산값 같음)은 차이로 세지 않음")
    ap.add_argument("--loose-text", action="store_true")
    ap.add_argument("--loose-formulas", action="store_true")
    ap.add_argument("--recalc", action="store_true", help="계산값이 비면 임시 사본을 엑셀로 다시 계산해 읽기")
    ap.add_argument("--max", type=int, default=40)
    ap.add_argument("--json", default="", help="전체 결과 JSON 저장 경로")
    a = ap.parse_args(argv)

    p1, p2 = Path(a.first), Path(a.second)
    for p in (p1, p2):
        if not p.exists():
            print(f"[오류] 파일이 없습니다: {p}", file=sys.stderr)
            return 2
    try:
        if a.period:
            res = period_compare(p1, p2, a)
            print_period(res, a)
            code = 0  # 기간 비교는 차이가 있는 것이 정상 - 알려 주기만 한다
        else:
            res = compare(p1, p2, a)
            if a.recalc:
                def _empty(side):
                    n, m = res["missing_calc"][side]
                    return n > 0 and m / n > 0.5
                if _empty("made") or _empty("ref"):
                    m_path = recalc_copy(p1) if _empty("made") else p1
                    r_path = recalc_copy(p2) if _empty("ref") else p2
                    res = compare(m_path, r_path, a)
                    res["made"], res["ref"] = str(p1), str(p2)
                    res["notes"].append("계산값이 비어 있던 쪽은 임시 사본을 엑셀로 다시 계산해서 비교했습니다(원본은 그대로).")
            print_result(res, a)
            code = 0 if not res["diffs"] else 1
    except SystemExit:
        raise
    except Exception as e:  # 사용자에게 보이는 오류는 짧게
        print(f"[오류] {type(e).__name__}: {e}", file=sys.stderr)
        return 2
    if a.json:
        Path(a.json).parent.mkdir(parents=True, exist_ok=True)
        Path(a.json).write_text(json.dumps(res, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
        print(f"\n전체 결과: {a.json}")
    return code


if __name__ == "__main__":
    sys.exit(main())
