"""작업 폴더 자료 목록 - 업무를 파악할 때 맨 먼저 돌린다.

파일마다 크기 · 수정일 · 이름 속 기간 표기(2026-06 · 26.2Q · 2분기 · 260814 …)를 뽑고,
기간만 다른 같은 이름의 파일을 '묶음'으로 모아 보여 준다(예: 대손충당금_결산파일_{기간} → 2026-03 · 2026-06).
묶음이 보이면 '지난 기간 완성본 · 그 전 기간 완성본 · 이번 기간 원자료'를 고르기 쉽다.
엑셀은 시트 이름과 크기, 워드 · PPT 는 문단 · 쪽 수를 함께 적는다.

사용: python -X utf8 inventory.py <폴더> [--depth 4] [--md 결과.md]
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path

SKIP_DIRS = {".git", "__pycache__", "_backup", "_xlsync", "node_modules", ".claude"}
SKIP_FILES = {"CLAUDE.md", "desktop.ini", "Thumbs.db"}
OFFICE = {".xlsx", ".xlsm", ".docx", ".pptx"}

# 기간 표기 - 앞의 것부터 찾는다(긴 표기 우선)
PERIOD_PATTERNS = [
    ("YYYY-MM-DD", re.compile(r"(?<!\d)(20\d{2})[-._](0[1-9]|1[0-2])[-._]([0-3]\d)(?!\d)")),
    ("YYYYMMDD", re.compile(r"(?<!\d)(20\d{2})(0[1-9]|1[0-2])([0-3]\d)(?!\d)")),
    ("YYYY-MM", re.compile(r"(?<!\d)(20\d{2})[-._](0[1-9]|1[0-2])(?!\d)")),
    ("YYYYMM", re.compile(r"(?<!\d)(20\d{2})(0[1-9]|1[0-2])(?!\d)")),
    ("FYyy.nQ", re.compile(r"FY\s?(\d{2})\.?\s?([1-4])Q", re.I)),
    ("yy.nQ", re.compile(r"(?<!\d)(\d{2})\.([1-4])Q", re.I)),
    ("YYYY nQ", re.compile(r"(?<!\d)(20\d{2})\s?[-.]?\s?([1-4])Q", re.I)),
    ("n분기", re.compile(r"([1-4])\s?분기")),
    ("nQ", re.compile(r"(?<![A-Za-z0-9])([1-4])Q(?![A-Za-z])", re.I)),
    ("반기", re.compile(r"(상반기|하반기|반기|1H|2H)")),
    ("YYMMDD", re.compile(r"(?<!\d)(2\d)(0[1-9]|1[0-2])([0-3]\d)(?!\d)")),
    ("YYMM", re.compile(r"(?<!\d)(2\d)(0[1-9]|1[0-2])(?!\d)")),
    ("YYYY년", re.compile(r"(20\d{2})\s?년")),
    ("FYYYYY", re.compile(r"FY\s?(20\d{2})", re.I)),
]
VERSION_RE = re.compile(r"([_\-\s(]?(v\d+(\.\d+)?|ver\d+|\d+차|최종|final|수정|rev\d*)[)\s]?)", re.I)


def find_periods(name: str) -> list[tuple[str, str, tuple[int, int]]]:
    """[(종류, 원문, (시작, 끝))] - 겹치지 않게."""
    taken: list[tuple[int, int]] = []
    out = []
    for kind, rx in PERIOD_PATTERNS:
        for m in rx.finditer(name):
            span = m.span()
            if any(not (span[1] <= a or span[0] >= b) for a, b in taken):
                continue
            taken.append(span)
            out.append((kind, m.group(0), span))
    out.sort(key=lambda x: x[2][0])
    return out


def series_key(name: str) -> str:
    """이름에서 기간 · 판(버전) 표기를 지운 묶음 열쇠."""
    stem, suffix = Path(name).stem, Path(name).suffix.lower()
    spans = [p[2] for p in find_periods(stem)]
    chars = list(stem)
    for a, b in sorted(spans, reverse=True):
        chars[a:b] = list("{기간}")
    key = "".join(chars)
    key = VERSION_RE.sub("", key)
    key = re.sub(r"[\s_\-]+", "_", key).strip("_")
    return key + suffix


def office_summary(p: Path) -> str:
    try:
        if p.suffix.lower() in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(p, read_only=True)
            parts = []
            for ws in wb.worksheets[:12]:
                try:
                    dim = ws.calculate_dimension() if ws.max_row else ""
                except Exception:
                    dim = ""
                parts.append(f"{ws.title}({dim})" if dim and dim != "A1:A1" else ws.title)
            more = f" 외 {len(wb.sheetnames) - 12}개" if len(wb.sheetnames) > 12 else ""
            wb.close()
            return f"시트 {len(wb.sheetnames)}: " + ", ".join(parts) + more
        if p.suffix.lower() == ".docx":
            with zipfile.ZipFile(p) as z:
                xml = z.read("word/document.xml").decode("utf8", "ignore")
            return f"문단 약 {xml.count('</w:p>')} · 표 {xml.count('</w:tbl>')}"
        if p.suffix.lower() == ".pptx":
            with zipfile.ZipFile(p) as z:
                n = len([x for x in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", x)])
            return f"슬라이드 {n}"
    except Exception as e:  # 깨진 파일 · 암호 걸린 파일
        return f"(못 읽음: {type(e).__name__})"
    return ""


def pdf_pages(p: Path) -> str:
    try:
        data = p.read_bytes()
        n = len(re.findall(rb"/Type\s*/Page[^s]", data))
        return f"약 {n}쪽" if n else ""
    except Exception:
        return ""


def human_size(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f}{unit}" if unit == "B" else f"{n:.1f}{unit}"
        n /= 1024
    return str(n)


def walk(root: Path, depth: int):
    root = root.resolve()
    for p in sorted(root.rglob("*")):
        rel = p.relative_to(root)
        if any(part in SKIP_DIRS for part in rel.parts):
            continue
        if len(rel.parts) > depth:
            continue
        if p.is_file() and not p.name.startswith("~$") and p.name not in SKIP_FILES:
            yield p, rel


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="작업 폴더 자료 목록(기간 표기 · 묶음 · 파일 요약)")
    ap.add_argument("folder")
    ap.add_argument("--depth", type=int, default=4)
    ap.add_argument("--md", default="", help="결과를 마크다운 파일로도 저장")
    ap.add_argument("--no-open", action="store_true", help="엑셀 · 워드 · PPT 안을 열지 않음(빠름)")
    a = ap.parse_args(argv)
    root = Path(a.folder)
    if not root.is_dir():
        print(f"[오류] 폴더가 없습니다: {root}", file=sys.stderr)
        return 2

    rows, series = [], defaultdict(list)
    locks = [p for p in root.rglob("~$*") if p.is_file()]
    for p, rel in walk(root, a.depth):
        st = p.stat()
        periods = find_periods(p.stem)
        ext = p.suffix.lower()
        info = ""
        if not a.no_open:
            if ext in OFFICE:
                info = office_summary(p)
            elif ext == ".pdf":
                info = pdf_pages(p)
        rows.append({
            "path": str(rel).replace("\\", "/"), "size": human_size(st.st_size),
            "mtime": dt.datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M"),
            "periods": " · ".join(x[1] for x in periods), "info": info,
        })
        if periods:
            series[(str(rel.parent).replace("\\", "/"), series_key(p.name))].append((periods[0][1], p.name))

    lines = [f"# 자료 목록 - {root.resolve()}", "", f"파일 {len(rows)}개 (깊이 {a.depth}, 숨김 · 백업 · .claude 제외)", ""]
    if locks:
        lines += ["**열려 있는 파일(잠금 ~$ 발견) - 작업 전에 닫아 달라고 말한다:** " +
                  ", ".join(str(x.relative_to(root)) for x in locks), ""]
    multi = {k: v for k, v in series.items() if len(v) >= 2}
    if multi:
        lines += ["## 기간 묶음 (이름이 같고 기간만 다른 파일)", "",
                  "| 폴더 | 묶음 | 기간(이름 순) |", "|---|---|---|"]
        for (folder, key), items in sorted(multi.items()):
            items.sort(key=lambda x: x[1])
            lines.append(f"| {folder} | `{key}` | " + " · ".join(i[0] for i in items) + " |")
        lines.append("")
    lines += ["## 파일", "", "| 파일 | 크기 | 수정 | 기간 표기 | 내용 |", "|---|---|---|---|---|"]
    for r in rows:
        lines.append(f"| {r['path']} | {r['size']} | {r['mtime']} | {r['periods']} | {r['info']} |")
    text = "\n".join(lines)
    print(text)
    if a.md:
        Path(a.md).parent.mkdir(parents=True, exist_ok=True)
        Path(a.md).write_text(text + "\n", encoding="utf-8")
        print(f"\n저장: {a.md}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
