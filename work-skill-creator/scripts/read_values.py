"""여러 엑셀 파일에서 칸 값을 한 번에 읽는다 - 업무 사이 대사 · 결과 보고 · 다음 기간 확인용.

주소로:   --ref "충당금계산!C31"
이름표로: --label "충당금계산:당기말 대손충당금 설정액"   (이름표 오른쪽 칸, --down 이면 아래 칸)
          --label "당기말 대손충당금 설정액"              (모든 시트에서 찾음)
계산값이 비어 있으면(openpyxl 로 저장한 파일) --recalc 로 임시 사본을 엑셀로 다시 계산해서 읽는다(원본은 그대로).

사용: python -X utf8 read_values.py <파일> [<파일> …] (--ref … | --label …)… [--offset 1] [--down] [--recalc] [--json 결과.json]
예(업무 사이 대사): read_values.py 대손_2026-06.xlsx 정산표_2026-06.xlsx --label "당기말 대손충당금 설정액" --label "대손충당금"
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))


def _fmt(v):
    import wsc_xl as W
    return W.fmt(v)


def read_one(path: Path, refs: list[str], labels: list[str], offset: int, down: bool) -> list[dict]:
    import openpyxl
    from openpyxl.utils import get_column_letter
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ref in refs:
        sheet, addr = ref.rsplit("!", 1)
        sheet = sheet.strip("'")
        if sheet not in wb.sheetnames:
            out.append({"file": path.name, "ask": ref, "cell": "", "value": None, "note": "시트 없음"})
            continue
        out.append({"file": path.name, "ask": ref, "cell": f"{sheet}!{addr.replace('$', '')}",
                    "value": wb[sheet][addr.replace("$", "")].value, "note": ""})
    for spec in labels:
        sheet, text = (spec.split(":", 1) if ":" in spec and spec.split(":", 1)[0] in wb.sheetnames else (None, spec))
        sheets = [wb[sheet]] if sheet else wb.worksheets
        found = False
        for ws in sheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.strip() == text.strip():
                        r, col = (c.row + offset, c.column) if down else (c.row, c.column + offset)
                        out.append({"file": path.name, "ask": spec, "cell": f"{ws.title}!{get_column_letter(col)}{r}",
                                    "value": ws.cell(r, col).value, "note": ""})
                        found = True
        if not found:
            out.append({"file": path.name, "ask": spec, "cell": "", "value": None, "note": "이름표 없음"})
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="여러 엑셀 파일에서 칸 값 읽기")
    ap.add_argument("files", nargs="+")
    ap.add_argument("--ref", action="append", default=[], help="시트!칸 (여러 번 가능)")
    ap.add_argument("--label", action="append", default=[], help="[시트:]이름표 (여러 번 가능)")
    ap.add_argument("--offset", type=int, default=1, help="이름표에서 몇 칸 떨어진 값(기본 1)")
    ap.add_argument("--down", action="store_true", help="이름표 아래 칸을 읽음(기본은 오른쪽)")
    ap.add_argument("--recalc", action="store_true", help="계산값이 비면 임시 사본을 엑셀로 다시 계산해 읽기")
    ap.add_argument("--json", default="")
    a = ap.parse_args(argv)
    if not a.ref and not a.label:
        print("[오류] --ref 나 --label 을 하나 이상 주세요", file=sys.stderr)
        return 2
    rows = []
    for f in a.files:
        p = Path(f)
        if not p.exists():
            rows.append({"file": p.name, "ask": "", "cell": "", "value": None, "note": "파일 없음"})
            continue
        got = read_one(p, a.ref, a.label, a.offset, a.down)
        if a.recalc and any(x["value"] is None and x["cell"] for x in got):
            from wsc_xl import excel_recalc
            tmp = Path(tempfile.mkdtemp(prefix="wsc_read_")) / p.name
            shutil.copy2(p, tmp)
            excel_recalc(tmp)
            got = read_one(tmp, a.ref, a.label, a.offset, a.down)
            for x in got:
                x["file"] = p.name
                x["note"] = (x["note"] + " 다시 계산해 읽음").strip()
        rows += got
    print("| 파일 | 찾은 것 | 칸 | 값 | 비고 |")
    print("|---|---|---|---|---|")
    for x in rows:
        print(f"| {x['file']} | {x['ask']} | {x['cell']} | {_fmt(x['value'])} | {x['note']} |")
    empty = [x for x in rows if x["value"] is None and x["cell"]]
    if empty and not a.recalc:
        print("\n[주의] 값이 빈 칸이 있습니다 - 계산값이 없는 파일이면 --recalc 를 주세요.")
    if a.json:
        Path(a.json).write_text(json.dumps(rows, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
