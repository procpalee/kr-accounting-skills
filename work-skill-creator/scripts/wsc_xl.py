"""업무 스킬 공용 도우미 - 업무 스킬의 scripts/ 에 복사해서 쓴다(외부 의존: openpyxl, 엑셀 재계산은 pywin32).

work-skill-creator 가 만드는 업무 스킬은 모두 같은 원칙을 지킨다.
  - 받은 자료(원본)는 읽기만 한다. 결과는 새 파일로 만들고, 이미 있으면 덮어쓰지 않는다.
  - 엑셀에 열려 있는 파일(~$ 잠금 파일)은 건드리지 않고 멈춘다.
  - 지난 기간 완성본을 사본으로 이어받아 바뀌는 칸만 고친다(서식 · 구조는 사람이 만든 그대로).
  - 계산은 엑셀 수식으로 남기고, 저장한 뒤 엑셀로 다시 계산한다(openpyxl 은 수식을 계산하지 않는다).
이 파일은 그 원칙을 매번 같은 코드로 지키기 위한 함수 모음이다.

주요 함수
  check_unlocked(*paths)            엑셀에 열려 있으면 멈춤
  guard_not_received(path)          받은 자료 폴더(01_받은자료) 파일을 제자리에서 고치려 하면 멈춤
  backup(path)                      _backup 폴더에 시각 붙인 사본
  copy_new(src, dst)                사본 만들기(대상이 있으면 멈춤)
  find_latest(dirs, keyword)        이름에 keyword 가 든 파일 중 최근 수정본 + 후보 목록
  excel_recalc(path)                엑셀(별도 인스턴스)로 열어 전체 다시 계산 후 저장
  ExcelEdit(path)                   그림 · 메모 · 배열수식이 있는 파일을 엑셀 프로그램으로 직접 고칠 때
  xls_to_xlsx(src, out_dir)         옛 형식 .xls 를 .xlsx 사본으로(원본은 그대로)
  cached_values(path, refs)         계산값 읽기(비어 있으면 임시 사본을 다시 계산해서 읽음)
  find_label(ws, text)              이름표(글자)로 칸 찾기 - 행이 늘고 줄어도 위치를 따라간다
  clear_row_heights(ws)             openpyxl 이 굳힌 고정 행 높이를 지워 엑셀 자동 높이로
  compare_values / print_compare    기대값과 결과값 대조
  round_half_up(x, n)               엑셀식 반올림(5 이상 올림)
  Period('2026-06')                 기간 계산(직전 기간, 분기 번호, 표기, 회계연도 결산월)
  shared_value(path, key, …)        여러 업무가 같이 쓰는 기간 확정값(환율 등) 읽기 · 확정
  cleanup_excel()                   이 도구가 띄웠다가 남은 숨은 엑셀만 정리 (명령: python wsc_xl.py --cleanup)
"""
from __future__ import annotations

import calendar
import datetime as dt
import gc
import json
import shutil
import sys
import tempfile
import time
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

SKIP_PARTS = {"_backup", "_xlsync", "__pycache__", ".git"}
RECEIVED_DIRS = ("01_받은자료",)  # 받은 자료(원본) 폴더 이름 - 작업 폴더 체계가 다르면 여기에 더한다


# ── 파일 안전 ────────────────────────────────────────────────
def lock_files(path: Path) -> list[Path]:
    """엑셀 · 워드가 파일을 열면 같은 폴더에 만드는 잠금 파일(~$…) 후보."""
    p = Path(path)
    names = {"~$" + p.name, "~$" + p.name[1:], "~$" + p.name[2:]}
    return [p.with_name(n) for n in names if n != "~$"]


def is_locked(path: Path) -> bool:
    return any(f.exists() for f in lock_files(path))


def check_unlocked(*paths: Path) -> None:
    """하나라도 엑셀(워드 · PPT)에 열려 있으면 멈춘다."""
    for p in paths:
        if p and Path(p).exists() and is_locked(Path(p)):
            raise SystemExit(f"[멈춤] 프로그램에 열려 있는 파일입니다(~$ 잠금 파일): {p}\n"
                             f"  파일을 닫은 뒤 다시 실행하세요. 엑셀 창이 안 보이면 지난 실행이 남긴 숨은 엑셀일 수 있습니다 -\n"
                             f"  작업 관리자로 엑셀을 끄지 말고(열어 둔 다른 파일까지 꺼짐) `python -X utf8 wsc_xl.py --cleanup` 으로 정리하세요.")


def guard_not_received(path: Path) -> None:
    """받은 자료 폴더 안 파일을 제자리에서 고치거나 저장하려 하면 멈춘다(원본 보호)."""
    parts = Path(path).resolve().parts
    if any(d in parts for d in RECEIVED_DIRS):
        raise SystemExit(f"[멈춤] 받은 자료(원본) 폴더의 파일은 고치지 않습니다: {path}\n"
                         f"  02_작업 에 사본을 만들어 그 사본에 하세요.")


def backup(path: Path, backup_dir: Path | None = None) -> Path | None:
    """고치기 전 사본을 남긴다. 기본 위치는 같은 폴더의 _backup."""
    p = Path(path)
    if not p.exists():
        return None
    if backup_dir is None:
        guard_not_received(p)  # 받은 자료 폴더 안에 _backup 을 만들지 않는다
    d = Path(backup_dir) if backup_dir else p.parent / "_backup"
    d.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    dst = d / f"{p.stem}_{stamp}{p.suffix}"
    shutil.copy2(p, dst)
    return dst


def copy_new(src: Path, dst: Path) -> Path:
    """src 를 dst 로 복사한다. dst 가 이미 있으면 멈춘다(사람이 손댄 파일일 수 있다)."""
    src, dst = Path(src), Path(dst)
    if not src.exists():
        raise SystemExit(f"[멈춤] 원본이 없습니다: {src}")
    if dst.exists():
        raise SystemExit(f"[멈춤] 이미 있는 파일입니다(덮어쓰지 않음): {dst}\n"
                         f"  다시 만들려면 기존 파일 이름을 바꾸거나 다른 출력 경로를 주세요.")
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dst)
    return dst


def find_latest(dirs, keyword: str, exts=(".xlsx", ".xlsm")) -> tuple[Path | None, list[Path]]:
    """dirs 아래(하위 폴더 포함)에서 이름에 keyword 가 든 파일 중 가장 최근 수정본과 후보 전체(최근순).

    '완성본'이 여러 판(v1 · v2 · 날짜)으로 남는 경우가 많다. 고른 파일과 후보를 계획에 보여 주고 사용자가 확인하게 한다.
    """
    cands = []
    for d in ([dirs] if isinstance(dirs, (str, Path)) else dirs):
        d = Path(d)
        if not d.is_dir():
            continue
        for p in d.rglob("*"):
            if not p.is_file() or p.suffix.lower() not in exts:
                continue
            if p.name.startswith("~$") or SKIP_PARTS & set(p.parts):
                continue
            if keyword in p.name:
                cands.append(p)
    cands.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return (cands[0] if cands else None), cands


# ── 엑셀 프로그램(COM) - 이 도구가 띄운 인스턴스만 다루고, 그것만 끈다 ───────────────
_REG = Path(tempfile.gettempdir()) / "wsc_excel_pids.json"


def _reg_load() -> list:
    try:
        return json.loads(_REG.read_text(encoding="utf-8"))
    except Exception:
        return []


def _reg_save(items: list) -> None:
    try:
        _REG.write_text(json.dumps(items), encoding="utf-8")
    except Exception:
        pass


def _proc_ctime(pid: int | None) -> str | None:
    """프로세스 생성 시각 - 기록한 PID 가 같은 프로세스인지(번호 재사용이 아닌지) 확인하는 데 쓴다."""
    if not pid:
        return None
    try:
        import win32api
        import win32process
        h = win32api.OpenProcess(0x1000, False, pid)  # PROCESS_QUERY_LIMITED_INFORMATION
        try:
            return str(win32process.GetProcessTimes(h)["CreationTime"])
        finally:
            win32api.CloseHandle(h)
    except Exception:
        return None


def _alive(pid: int | None) -> bool:
    if not pid:
        return False
    try:
        import win32api
        import win32process
        h = win32api.OpenProcess(0x1000, False, pid)
        try:
            return win32process.GetExitCodeProcess(h) == 259  # STILL_ACTIVE
        finally:
            win32api.CloseHandle(h)
    except Exception:
        return False


def _terminate(pid: int | None, ctime: str | None) -> bool:
    """기록한 생성 시각과 같을 때만(같은 프로세스일 때만) 끈다 - 사용자 엑셀을 건드리지 않기 위해."""
    if not pid or not ctime or _proc_ctime(pid) != ctime:
        return False
    try:
        import win32api
        h = win32api.OpenProcess(0x0001, False, pid)  # PROCESS_TERMINATE
        try:
            win32api.TerminateProcess(h, 1)
            return True
        finally:
            win32api.CloseHandle(h)
    except Exception:
        return False


class _Excel:
    """새 엑셀 인스턴스 하나(사용자가 열어 둔 엑셀과 섞이지 않게 DispatchEx). quit() 이 그 인스턴스만 끈다."""

    def __init__(self):
        try:
            import pythoncom
            import win32com.client
        except ImportError as e:  # pragma: no cover - 환경에 따라
            raise RuntimeError("엑셀 재계산에는 pywin32 가 필요합니다(python -m pip install pywin32). "
                               "윈도우 + 데스크톱 엑셀이 있어야 합니다.") from e
        self._pythoncom = pythoncom
        pythoncom.CoInitialize()
        self.xl = None
        self.pid = self.ctime = None
        try:
            self.xl = win32com.client.DispatchEx("Excel.Application")
            self.xl.Visible = False
            self.xl.DisplayAlerts = False
            self.xl.EnableEvents = False
            try:
                self.xl.AutomationSecurity = 3  # 매크로 실행 안 함
            except Exception:
                pass
            try:
                import win32process
                self.pid = win32process.GetWindowThreadProcessId(self.xl.Hwnd)[1]
                self.ctime = _proc_ctime(self.pid)
            except Exception:
                self.pid = self.ctime = None
            if self.pid:
                items = _reg_load()
                items.append({"pid": self.pid, "ctime": self.ctime, "started": dt.datetime.now().isoformat(timespec="seconds")})
                _reg_save(items)
        except Exception:
            self.quit()
            raise

    def quit(self) -> None:
        try:
            if self.xl is not None:
                try:
                    self.xl.Quit()
                except Exception:
                    pass
        finally:
            self.xl = None
            gc.collect()
            if self.pid:
                for _ in range(25):  # 5초 기다려도 남아 있으면 이 인스턴스만 끈다
                    if not _alive(self.pid):
                        break
                    time.sleep(0.2)
                else:
                    _terminate(self.pid, self.ctime)
                _reg_save([x for x in _reg_load() if x.get("pid") != self.pid])
            try:
                self._pythoncom.CoUninitialize()
            except Exception:
                pass


def cleanup_excel(min_age_minutes: float = 15, everything: bool = False) -> list[int]:
    """이 도구가 띄웠다가 남은 숨은 엑셀만 끈다(기록한 PID + 생성 시각이 같은 것만).

    실행 중인 다른 작업의 엑셀을 끄지 않도록 기본은 15분이 지난 기록만 정리한다(everything=True 면 나이 무시).
    """
    killed, keep = [], []
    now = dt.datetime.now()
    for it in _reg_load():
        pid, ctime = it.get("pid"), it.get("ctime")
        try:
            age = (now - dt.datetime.fromisoformat(it.get("started"))).total_seconds() / 60
        except Exception:
            age = 10 ** 6
        if not _alive(pid) or _proc_ctime(pid) != ctime:
            continue  # 이미 끝났거나 다른 프로세스 - 기록만 지운다
        if not everything and age < min_age_minutes:
            keep.append(it)
            continue
        if _terminate(pid, ctime):
            killed.append(pid)
        else:
            keep.append(it)
    _reg_save(keep)
    return killed


def excel_recalc(path: Path, save_to: Path | None = None) -> Path:
    """엑셀로 열어 전체 다시 계산하고 저장한다. save_to 를 주면 그 경로에 사본을 만들어 계산한다.

    openpyxl 로 쓴 수식 칸은 계산값이 비어 있다. 결과를 읽거나 비교하기 전에 반드시 한 번 거친다.
    엑셀은 상대 경로를 못 찾으므로 항상 절대 경로로 연다. 외부 연결은 새로 고치지 않는다(UpdateLinks=0).
    """
    src = Path(path).resolve()
    dst = Path(save_to).resolve() if save_to else src
    if save_to:
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
    else:
        guard_not_received(src)
    ex = _Excel()
    try:
        wb = ex.xl.Workbooks.Open(str(dst), 0, False)
        ex.xl.CalculateFull()
        wb.Save()
        wb.Close(False)
        wb = None
    finally:
        ex.quit()
    return dst


class ExcelEdit:
    """엑셀 프로그램(COM)으로 파일을 열어 고치고, 정상 종료하면 다시 계산해 저장한다.

    그림 · 메모 · 스레드 메모 · 웹 추가기능 · 슬라이서 · 배열수식이 있는 파일은 openpyxl 로 저장하면 그 요소가 사라지거나
    엑셀이 파일을 못 여는 일이 생긴다. 그런 파일은 이 클래스로 고친다.
        with ExcelEdit(path) as (xl, wb):
            ws = wb.Worksheets("시트")
            xl_put(ws, "C2", "=DATE(2026,9,30)")
    행을 넣고 빼면 다른 시트의 참조 범위도 엑셀이 따라 움직여 준다(openpyxl 은 안 움직인다).
    """

    def __init__(self, path: Path):
        self.path = Path(path).resolve()
        self.ex = self.wb = None

    def __enter__(self):
        guard_not_received(self.path)
        self.ex = _Excel()
        try:
            self.wb = self.ex.xl.Workbooks.Open(str(self.path), 0, False)
        except Exception:
            self.ex.quit()  # 열다 실패해도 숨은 엑셀을 남기지 않는다
            raise
        return self.ex.xl, self.wb

    def __exit__(self, et, ev, tb):
        try:
            if et is None:
                self.ex.xl.CalculateFull()
                self.wb.Save()
            self.wb.Close(False)
        finally:
            self.wb = None
            self.ex.quit()
        return False


def xl_put(ws, addr: str, v) -> None:
    """COM 시트에 값 또는 수식('=' 로 시작하는 글)을 쓴다. None 이면 내용만 지운다."""
    rng = ws.Range(addr)
    if v is None:
        rng.ClearContents()
    elif isinstance(v, str) and v.startswith("="):
        rng.Formula = v
    else:
        if isinstance(v, dt.date) and not isinstance(v, dt.datetime):
            v = dt.datetime(v.year, v.month, v.day)
        rng.Value = v


def xls_to_xlsx(src: Path, out_dir: Path) -> Path:
    """옛 형식 .xls(더존 등 내려받기)를 .xlsx 사본으로 만든다. 원본은 임시 사본으로만 연다(받은 자료 폴더에 잠금 파일이 생기지 않게)."""
    src = Path(src).resolve()
    out_dir = Path(out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    dst = out_dir / (src.stem + ".xlsx")
    if dst.exists():
        raise SystemExit(f"[멈춤] 이미 있는 파일입니다(덮어쓰지 않음): {dst}")
    with tempfile.TemporaryDirectory(prefix="wsc_xls_", ignore_cleanup_errors=True) as td:
        tmp = Path(td) / src.name
        shutil.copy2(src, tmp)
        ex = _Excel()
        try:
            wb = ex.xl.Workbooks.Open(str(tmp), 0, True)
            wb.SaveAs(str(dst), 51)  # 51 = .xlsx
            wb.Close(False)
            wb = None
        finally:
            ex.quit()
    return dst


# ── 읽기 ────────────────────────────────────────────────────
def _split_ref(ref: str) -> tuple[str, str]:
    sheet, addr = ref.rsplit("!", 1)
    return sheet.strip("'"), addr.replace("$", "")


def cached_values(path: Path, refs: list[str], recalc_if_empty: bool = True) -> dict:
    """'시트!A1' 목록의 계산값. 하나라도 비어 있으면 임시 사본을 엑셀로 다시 계산해서 읽는다(원본은 그대로, 사본은 지운다)."""
    import openpyxl

    def _read(p):
        wb = openpyxl.load_workbook(p, data_only=True)
        out = {}
        for r in refs:
            s, a = _split_ref(r)
            out[r] = wb[s][a].value
        return out

    vals = _read(path)
    if recalc_if_empty and any(v is None for v in vals.values()):
        with tempfile.TemporaryDirectory(prefix="wsc_", ignore_cleanup_errors=True) as td:
            tmp = Path(td) / Path(path).name
            excel_recalc(Path(path), tmp)
            vals = _read(tmp)
    return vals


def find_label(ws, text: str, *, exact: bool = True, cols=None, rows=None):
    """이름표(글자)가 든 첫 칸의 (행, 열). 없으면 None. (openpyxl 시트용)

    행 번호를 코드에 박아 두면 다음 기간에 행이 늘거나 줄 때 틀린 칸을 읽는다.
    '합계' · '1. 차입금' 같은 이름표로 찾으면 위치가 바뀌어도 따라간다.
    cols / rows 로 찾는 범위를 좁힐 수 있다(예: cols=range(1, 4)).
    """
    want = str(text).strip()
    for row in ws.iter_rows():
        for c in row:
            if cols is not None and c.column not in cols:
                continue
            if rows is not None and c.row not in rows:
                continue
            v = c.value
            if not isinstance(v, str):
                continue
            s = v.strip()
            if (s == want) if exact else (want in s):
                return c.row, c.column
    return None


def clear_row_heights(ws, rows=None) -> int:
    """고정된 행 높이를 지워 엑셀 자동 높이로 되돌린다(openpyxl 시트). 지운 행 수를 돌려준다.

    엑셀이 저장한 자동 높이는 openpyxl 로 다시 저장하면 고정 높이(customHeight)로 굳고,
    copy_worksheet · insert_rows 를 거치면 엉뚱한 행에 붙는다. 사람이 일부러 고정한 높이가 없는 시트라면
    저장 직전에 이 함수로 지운다(지난 완성본에 고정 높이가 있었는지 먼저 compare_xlsx --formats 로 확인).
    """
    n = 0
    for r, d in list(ws.row_dimensions.items()):  # openpyxl 에서는 높이 값(ht)이 있으면 곧 '고정'이다(customHeight 는 ht 로 정해짐)
        if rows is not None and r not in rows:
            continue
        if d.ht is not None:
            d.ht = None
            n += 1
    return n


def find_label_all(ws, text: str, *, exact: bool = True) -> list[tuple[int, int]]:
    want = str(text).strip()
    hits = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str):
                s = c.value.strip()
                if (s == want) if exact else (want in s):
                    hits.append((c.row, c.column))
    return hits


# ── 공유 확정값 ──────────────────────────────────────────────
def shared_value(path: Path, key: str, value=None, source: str = "", force: bool = False):
    """여러 업무가 같이 쓰는 기간 값(기말 환율 등)을 한 파일에서 읽거나 확정한다.

    value 가 None 이면 읽기만(없으면 None). 값을 주면 저장하되, 이미 다른 값으로 확정돼 있으면 멈춘다(바꾸려면 force=True).
    """
    p = Path(path)
    data = json.loads(p.read_text(encoding="utf-8-sig")) if p.exists() else {}
    cur = data.get(key)
    if value is None:
        return cur.get("값") if isinstance(cur, dict) else None
    if isinstance(cur, dict) and cur.get("값") != value and not force:
        raise SystemExit(f"[멈춤] '{key}' 는 이미 {cur.get('값')} 로 확정돼 있습니다({cur.get('확정일')}, {cur.get('출처')}). "
                         f"새 값 {value} 로 바꾸려면 사용자 확인 뒤 force 로 다시 저장하세요.")
    data[key] = {"값": value, "출처": source, "확정일": dt.date.today().isoformat()}
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    return value


# ── 대조 ────────────────────────────────────────────────────
def round_half_up(x, ndigits: int = 0):
    """엑셀 ROUND 와 같은 반올림(0.5 는 올림). 파이썬 round 는 짝수 쪽으로 보내서 1원씩 어긋난다."""
    if x is None:
        return None
    q = Decimal(1).scaleb(-ndigits)
    r = Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP)
    return int(r) if ndigits <= 0 else float(r)


def compare_values(actual: dict, expected: dict, tol: float = 0.5, blank_as_zero: bool = True) -> list[tuple]:
    """[(키, 기대, 결과, 차이, 일치여부)]. 숫자는 tol 이내면 일치."""
    rows = []
    for k, e in expected.items():
        a = actual.get(k)
        if blank_as_zero:
            a = 0 if a is None else a
            e = 0 if e is None else e
        if isinstance(e, (int, float)) and isinstance(a, (int, float)) and not isinstance(e, bool):
            d = a - e
            rows.append((k, e, a, d, abs(d) <= tol))
        else:
            rows.append((k, e, a, None, e == a))
    return rows


def fmt(v) -> str:
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, int):
        return f"{v:,}"
    if isinstance(v, float):
        return f"{v:,.2f}"
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return "" if v is None else str(v)


def print_compare(actual: dict, expected: dict, label: str, tol: float = 0.5) -> int:
    """대조 결과를 출력하고 불일치 수를 돌려준다."""
    rows = compare_values(actual, expected, tol)
    bad = [r for r in rows if not r[4]]
    print(f"\n[대조] {label} - {len(rows) - len(bad)}/{len(rows)} 일치")
    for k, ev, av, d, _ in bad:
        extra = f" · 차이 {fmt(d)}" if d is not None else ""
        print(f"  ✗ {k}: 기대 {fmt(ev)} · 결과 {fmt(av)}{extra}")
    return len(bad)


# ── 기간 ────────────────────────────────────────────────────
class Period:
    """결산 기간 하나. Period('2026-06') · Period('2026-06-30') · Period('202606') · Period('2026.06').

    kind: 'Q' 분기(기본) · 'M' 월 · 'H' 반기 · 'Y' 연.  fy_end: 회계연도 마지막 달(기본 12, 3월 결산이면 3).
    """

    STEP = {"M": 1, "Q": 3, "H": 6, "Y": 12}

    def __init__(self, s, kind: str = "Q", fy_end: int = 12):
        if isinstance(s, Period):
            self.year, self.month = s.year, s.month
        elif isinstance(s, (dt.date, dt.datetime)):
            self.year, self.month = s.year, s.month
        else:
            digits = "".join(ch for ch in str(s) if ch.isdigit())
            if len(digits) < 6:
                raise ValueError(f"기간을 읽을 수 없습니다: {s!r} (예: 2026-06)")
            self.year, self.month = int(digits[:4]), int(digits[4:6])
        if not 1 <= self.month <= 12:
            raise ValueError(f"월이 이상합니다: {s!r}")
        kind = kind.upper()
        if kind not in self.STEP:
            raise ValueError(f"kind 는 M · Q · H · Y 중 하나: {kind!r}")
        if not 1 <= int(fy_end) <= 12:
            raise ValueError(f"fy_end(회계연도 마지막 달)가 이상합니다: {fy_end!r}")
        self.kind, self.fy_end = kind, int(fy_end)
        if (self.month - self.fy_end) % self.STEP[kind]:
            raise ValueError(f"{s!r} 는 {kind} 기간의 마지막 달이 아닙니다(회계연도 마지막 달 {self.fy_end}월 기준)")

    def __repr__(self):
        extra = f", fy_end={self.fy_end}" if self.fy_end != 12 else ""
        return f"Period('{self.ym}', '{self.kind}'{extra})"

    def __eq__(self, other):
        return isinstance(other, Period) and (self.year, self.month, self.kind, self.fy_end) == \
            (other.year, other.month, other.kind, other.fy_end)

    def __hash__(self):
        return hash((self.year, self.month, self.kind, self.fy_end))

    @property
    def end(self) -> dt.date:
        return dt.date(self.year, self.month, calendar.monthrange(self.year, self.month)[1])

    @property
    def q(self) -> int:
        """회계연도 기준 분기 번호(12월 결산이면 달력 분기와 같다)."""
        return ((self.month - self.fy_end - 1) % 12) // 3 + 1

    @property
    def fy_year(self) -> int:
        """회계연도(그 회계연도가 끝나는 해)."""
        return self.year if self.month <= self.fy_end else self.year + 1

    @property
    def ym(self) -> str:
        return f"{self.year}-{self.month:02d}"

    @property
    def dot(self) -> str:
        return f"{self.year}.{self.month:02d}"

    @property
    def yq(self) -> str:
        return f"{self.fy_year}-{self.q}Q"

    @property
    def fy(self) -> str:
        return f"FY{self.fy_year % 100:02d}.{self.q}Q"

    @property
    def kor(self) -> str:
        yr = f"{self.year}년" if self.fy_end == 12 else f"{self.fy_year} 회계연도"
        half = "상" if self.q <= 2 else "하"
        return {"M": f"{self.year}년 {self.month}월", "Q": f"{yr} {self.q}분기",
                "H": f"{yr} {half}반기", "Y": yr}[self.kind]

    def shift(self, n: int) -> "Period":
        m = self.year * 12 + (self.month - 1) + n * self.STEP[self.kind]
        return Period(f"{m // 12}-{m % 12 + 1:02d}", self.kind, self.fy_end)

    def months(self) -> list["Period"]:
        """이 기간 안의 달(월 기간)들 - 분기면 석 달. 기간 중 들어온 계약서처럼 달마다 오는 자료를 찾을 때."""
        n = self.STEP[self.kind]
        return [Period(self, "M").shift(-k) for k in range(n - 1, -1, -1)]

    def prev(self) -> "Period":
        return self.shift(-1)

    def next(self) -> "Period":
        return self.shift(1)

    @property
    def is_year_start(self) -> bool:
        """회계연도 첫 기간(12월 결산이면 1월 · 1분기 · 상반기)인가 - 연초 이월 · 연간 갱신을 따로 처리해야 하는 때."""
        return (self.month - self.fy_end - self.STEP[self.kind]) % 12 == 0


# ── 명령 ────────────────────────────────────────────────────
if __name__ == "__main__":
    import argparse

    ap = argparse.ArgumentParser(description="work-skill-creator 공용 도우미")
    ap.add_argument("--cleanup", action="store_true", help="이 도구가 띄웠다가 남은 숨은 엑셀만 정리")
    ap.add_argument("--all", action="store_true", help="--cleanup 때 15분 안 된 기록도 정리(다른 작업이 돌고 있지 않을 때만)")
    a = ap.parse_args()
    if a.cleanup:
        killed = cleanup_excel(everything=a.all)
        print(f"정리한 엑셀: {len(killed)}개" + (f" (PID {', '.join(map(str, killed))})" if killed else "")
              + " - 이 도구가 띄운 것만 정리했습니다. 사용자가 연 엑셀은 그대로입니다.")
        left = _reg_load()
        if left:
            print(f"아직 돌고 있는 기록 {len(left)}개(15분 안 됨) - 다른 작업이 쓰는 중일 수 있어 그대로 뒀습니다. 확실하면 --all.")
        sys.exit(0)
    ap.print_help()
